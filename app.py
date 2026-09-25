import os
import json
import urllib.parse
from datetime import datetime, timedelta
from io import BytesIO
from fastapi import FastAPI, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from supabase import create_client, Client
from dotenv import load_dotenv
from auth import (
    create_session, get_current_user, COOKIE_NAME,
    verify_login, get_user_info, is_admin, log_activity
)
from barcode_routes import router as barcode_router

load_dotenv()

SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    raise Exception("Missing SUPABASE_URL or SUPABASE_KEY")

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

app = FastAPI(title="Building Materials Shop")
app.include_router(barcode_router)

SHOP_NAME = "OBOLO TILES & CEMENT"
SHOP_PHONE = "053500108"
SHOP_ADDRESS = "MENZEZOR, GHANA"


def page(title, body, user=None, role=None, extra_head=""):
    if user and role == "admin":
        menu = ("<a href='/'>Home</a>"
                "<a href='/products'>Materials</a>"
                "<a href='/scan'>📷 Scan</a>"
                "<a href='/barcodes'>🏷️ Barcodes</a>"
                "<a href='/customers'>Customers</a>"
                "<a href='/suppliers'>Suppliers</a>"
                "<a href='/add'>Add</a>"
                "<a href='/sell'>New Sale</a>"
                "<a href='/cart'>Cart</a>"
                "<a href='/expenses'>💰 Expenses</a>"
                "<a href='/margins'>📈 Margins</a>"
                "<a href='/pnl'>📊 P&L</a>"
                "<a href='/activity'>📝 Activity</a>"
                "<a href='/alerts'>🚨 Alerts</a>"
                "<a href='/summary'>📅 Daily</a>"
                "<a href='/reports'>Reports</a>"
                "<a href='/users'>Users</a>"
                "<a href='/account'>👤 My Account</a>")
    elif user and role == "cashier":
        menu = ("<a href='/'>Home</a>"
                "<a href='/products'>Materials</a>"
                "<a href='/scan'>📷 Scan</a>"
                "<a href='/customers'>Customers</a>"
                "<a href='/suppliers'>Suppliers</a>"
                "<a href='/sell'>New Sale</a>"
                "<a href='/cart'>Cart</a>"
                "<a href='/account'>👤 My Account</a>")
    else:
        menu = ""

    user_bar = ""
    if user:
        role_display = f" ({role})" if role else ""
        user_bar = f'<span class="user-bar">👤 {user}{role_display} · <a href="/logout">Logout</a></span>'

    return f"""<!DOCTYPE html>
<html>
<head>
<title>{title} - {SHOP_NAME}</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
{extra_head}
<style>
* {{ box-sizing: border-box; }}
body {{ font-family: Arial, sans-serif; margin: 0; background: #f4f4f7; color: #222; }}
.header {{ background: #1e40af; color: white; padding: 15px 20px; }}
.header-top {{ display: flex; justify-content: space-between; align-items: center; }}
.header h1 {{ margin: 0; font-size: 20px; }}
.menu-toggle {{ display: none; background: transparent; border: 2px solid white; color: white; font-size: 22px; padding: 5px 12px; border-radius: 5px; cursor: pointer; }}
.menu-links {{ display: flex; flex-wrap: wrap; align-items: center; margin-top: 8px; }}
.menu-links a {{ color: white; text-decoration: none; margin-right: 15px; font-size: 14px; padding: 4px 0; }}
.menu-links a:hover {{ text-decoration: underline; }}
.user-bar {{ color: white; font-size: 13px; margin-left: auto; }}
.user-bar a {{ color: white; text-decoration: underline; }}
.container {{ max-width: 1100px; margin: 20px auto; padding: 0 15px; }}
.card {{ background: white; border-radius: 8px; padding: 20px; margin-bottom: 15px; box-shadow: 0 1px 3px rgba(0,0,0,0.06); }}
h2 {{ color: #1e40af; margin-top: 0; }}
h3 {{ color: #1e40af; margin-top: 15px; }}
.table-wrap {{ overflow-x: auto; -webkit-overflow-scrolling: touch; }}
table {{ width: 100%; border-collapse: collapse; min-width: 500px; }}
th, td {{ padding: 10px; text-align: left; border-bottom: 1px solid #eee; }}
th {{ background: #f9fafb; }}
input {{ width: 100%; padding: 8px 10px; margin: 4px 0 10px; border: 1px solid #ddd; border-radius: 5px; font-size: 14px; background: white; }}
select {{ width: 100%; padding: 6px 8px; margin: 4px 0 10px; border: 1px solid #ddd; border-radius: 5px; font-size: 14px; background: white; line-height: 1.2; }}
label {{ font-size: 14px; color: #333; font-weight: bold; }}
button, .btn {{ background: #1e40af; color: white; padding: 10px 16px; border: none; border-radius: 5px; cursor: pointer; font-size: 14px; text-decoration: none; display: inline-block; margin: 3px 0; }}
button:hover, .btn:hover {{ background: #1e3a8a; }}
.btn-success {{ background: #16a34a; }}
.btn-danger {{ background: #dc2626; }}
.btn-warn {{ background: #d97706; }}
.btn-small {{ padding: 6px 12px; font-size: 13px; }}
.btn-quick {{ background: #6b7280; padding: 6px 12px; font-size: 13px; margin: 2px; }}
.btn-quick:hover {{ background: #4b5563; }}
.low {{ color: #dc2626; font-weight: bold; }}
.ok {{ color: #16a34a; font-weight: bold; }}
.owed {{ color: #dc2626; font-weight: bold; font-size: 16px; }}
.clear {{ color: #16a34a; font-weight: bold; }}
.good-margin {{ color: #16a34a; font-weight: bold; }}
.ok-margin {{ color: #d97706; font-weight: bold; }}
.bad-margin {{ color: #dc2626; font-weight: bold; }}
.grid {{ display: grid; grid-template-columns: 1fr 1fr; gap: 15px; }}
.stat {{ text-align: center; padding: 15px; }}
.stat .num {{ font-size: 28px; font-weight: bold; color: #1e40af; }}
.stat .label {{ color: #666; font-size: 13px; }}
.login-box {{ max-width: 400px; margin: 80px auto; }}
.cart-total {{ background: #fef3c7; padding: 15px; border-radius: 8px; margin-top: 10px; font-size: 18px; }}
.alert-box {{ background: #fee2e2; border: 2px solid #dc2626; }}
.big-num {{ font-size: 42px; font-weight: bold; color: #1e40af; text-align: center; padding: 20px; }}
.profit-box {{ background: #dcfce7; border-left: 6px solid #16a34a; }}
.loss-box {{ background: #fee2e2; border-left: 6px solid #dc2626; }}
.success-msg {{ background: #dcfce7; border: 2px solid #16a34a; color: #166534; padding: 15px; border-radius: 8px; margin-bottom: 15px; }}
.error-msg {{ background: #fee2e2; border: 2px solid #dc2626; color: #991b1b; padding: 15px; border-radius: 8px; margin-bottom: 15px; }}
.date-bar {{ background: #dbeafe; border: 2px solid #1e40af; padding: 15px; border-radius: 8px; }}
.date-bar form {{ display: flex; gap: 10px; flex-wrap: wrap; align-items: flex-end; }}
.date-bar .field {{ flex: 1; min-width: 150px; }}
.period-badge {{ display: inline-block; background: #1e40af; color: white; padding: 6px 14px; border-radius: 20px; font-size: 14px; font-weight: bold; margin-bottom: 10px; }}

.pnl-table {{ width: 100%; border-collapse: collapse; font-size: 15px; }}
.pnl-table td {{ padding: 10px 15px; border-bottom: 1px solid #eee; }}
.pnl-table .label-col {{ text-align: left; }}
.pnl-table .amount-col {{ text-align: right; font-weight: bold; width: 180px; }}
.pnl-section-header {{ background: #1e40af; color: white; font-weight: bold; }}
.pnl-section-header td {{ padding: 10px 15px; }}
.pnl-subtotal {{ background: #f0f4ff; font-weight: bold; font-size: 16px; }}
.pnl-subtotal td {{ padding: 12px 15px; border-top: 2px solid #1e40af; border-bottom: 2px solid #1e40af; }}
.pnl-final {{ font-size: 18px; }}
.pnl-final td {{ padding: 15px; font-weight: bold; border-top: 3px double #1e40af; border-bottom: 3px double #1e40af; }}
.pnl-profit {{ background: #dcfce7; color: #166534; }}
.pnl-loss {{ background: #fee2e2; color: #991b1b; }}

.badge {{ display: inline-block; padding: 3px 10px; border-radius: 12px; font-size: 12px; font-weight: bold; }}
.badge-info {{ background: #dbeafe; color: #1e40af; }}
.badge-warning {{ background: #fef3c7; color: #92400e; }}
.badge-danger {{ background: #fee2e2; color: #991b1b; }}

.chart-container {{ position: relative; height: 300px; margin-top: 10px; }}

#reader {{ width: 100%; max-width: 500px; margin: 0 auto; border-radius: 8px; overflow: hidden; }}
.scan-result {{ padding: 15px; border-radius: 8px; margin-top: 15px; font-size: 16px; text-align: center; }}
.scan-success {{ background: #dcfce7; border: 2px solid #16a34a; color: #166534; }}
.scan-error {{ background: #fee2e2; border: 2px solid #dc2626; color: #991b1b; }}
.barcode-label {{ display: inline-block; padding: 10px; border: 1px solid #000; margin: 5px; background: white; text-align: center; width: 220px; }}
.barcode-label .name {{ font-weight: bold; font-size: 13px; margin-bottom: 2px; }}
.barcode-label .price {{ font-size: 12px; margin-bottom: 5px; }}
.barcode-label svg {{ max-width: 100%; }}
@media print {{
    .no-print, .header, .btn, button, .menu-links, .user-bar {{ display: none !important; }}
    .barcode-label {{ page-break-inside: avoid; }}
    body {{ background: white; }}
    .card {{ box-shadow: none; padding: 0; }}
}}

@media (max-width: 768px) {{
    .menu-toggle {{ display: block; }}
    .menu-links {{ display: none; flex-direction: column; align-items: stretch; margin-top: 12px; padding-top: 12px; border-top: 1px solid rgba(255,255,255,0.25); }}
    .menu-links.open {{ display: flex; }}
    .menu-links a {{ display: block; padding: 14px 5px; border-bottom: 1px solid rgba(255,255,255,0.15); margin: 0; font-size: 16px; }}
    .user-bar {{ display: block; margin: 12px 0 0 0; padding-top: 12px; border-top: 1px solid rgba(255,255,255,0.25); font-size: 14px; }}
    .container {{ margin: 10px auto; padding: 0 10px; }}
    .card {{ padding: 15px; }}
    .grid {{ grid-template-columns: 1fr; }}
    .stat .num {{ font-size: 24px; }}
    h2 {{ font-size: 20px; }}
    table {{ min-width: 450px; }}
    th, td {{ padding: 8px; font-size: 14px; }}
    button, .btn {{ width: 100%; text-align: center; margin: 5px 0; }}
    .btn-small {{ width: auto; }}
    .big-num {{ font-size: 32px; }}
    .date-bar form {{ flex-direction: column; }}
    .date-bar .field {{ width: 100%; }}
    .pnl-table td {{ padding: 8px 10px; font-size: 14px; }}
    .pnl-table .amount-col {{ width: 120px; }}
    .chart-container {{ height: 250px; }}
    .barcode-label {{ width: 100%; max-width: 220px; }}
}}
</style>
</head>
<body>
<div class="header">
<div class="header-top">
<h1>🏗️ {SHOP_NAME}</h1>
<button class="menu-toggle" onclick="document.getElementById('menuLinks').classList.toggle('open')">☰</button>
</div>
<div class="menu-links" id="menuLinks">
{menu}
{user_bar}
</div>
</div>
<div class="container">{body}</div>
</body>
</html>"""


# ============ HELPERS ============

def parse_dt(s):
    try:
        return datetime.fromisoformat(str(s).replace("Z", "+00:00")).replace(tzinfo=None)
    except Exception:
        return None


def log(username, action_type, description, severity="info", details=None):
    log_activity(username, action_type, description, severity, details)


# ============ AUTH ============

@app.get("/login", response_class=HTMLResponse)
def login_page(error: str = ""):
    err = f'<p style="color:#dc2626;">{error}</p>' if error else ""
    body = f"""
    <div class="login-box">
        <div class="card">
            <h2>🔐 Login</h2>
            {err}
            <form method="post" action="/login">
                <label>Username</label>
                <input type="text" name="username" required autofocus>
                <label>Password</label>
                <input type="password" name="password" required>
                <button type="submit" style="width:100%;">Login</button>
            </form>
        </div>
    </div>
    """
    return HTMLResponse(content=page("Login", body))


@app.post("/login")
async def do_login(username: str = Form(...), password: str = Form(...)):
    user_info = verify_login(username, password)
    if user_info:
        response = RedirectResponse("/", status_code=303)
        response.set_cookie(COOKIE_NAME, create_session(username), httponly=True, max_age=60*60*24*7)
        return response
    return RedirectResponse("/login?error=Wrong+username+or+password", status_code=303)


@app.get("/logout")
def logout(request: Request):
    username = get_current_user(request)
    if username:
        log(username, "logout", f"User '{username}' logged out")
    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie(COOKIE_NAME)
    return response


# ============ MY ACCOUNT ============

@app.get("/account", response_class=HTMLResponse)
def my_account(request: Request, msg: str = "", error: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    success_box = f'<div class="success-msg">✅ {msg}</div>' if msg else ""
    error_box = f'<div class="error-msg">❌ {error}</div>' if error else ""

    body = f"""
    <h2>👤 My Account</h2>
    {success_box}
    {error_box}

    <div class="card">
        <h3>Account Information</h3>
        <p><strong>Username:</strong> {username}</p>
        <p><strong>Full Name:</strong> {info.get('full_name','—') if info else '—'}</p>
        <p><strong>Role:</strong> {'Admin' if role == 'admin' else 'Cashier'}</p>
        <p><strong>Status:</strong> {'✅ Active' if info and info.get('is_active') else '❌ Inactive'}</p>
    </div>

    <div class="card">
        <h3>🔐 Change Password</h3>
        <form method="post" action="/account/change-password">
            <label>Current Password</label>
            <input type="password" name="current_password" required autocomplete="current-password">

            <label>New Password (min 4 characters)</label>
            <input type="password" name="new_password" required minlength="4" autocomplete="new-password">

            <label>Confirm New Password</label>
            <input type="password" name="confirm_password" required minlength="4" autocomplete="new-password">

            <button type="submit" class="btn btn-success">🔐 Change Password</button>
            <a href="/" class="btn">Cancel</a>
        </form>
    </div>

    <div class="card">
        <h3>ℹ️ Tips</h3>
        <ul>
            <li>Choose a password that is hard for others to guess</li>
            <li>Don't use your name, phone number, or birthday</li>
            <li>After changing your password, you'll be logged out and must log in again with the new password</li>
            <li>If you forget your password, ask the admin to reset it from the Users page</li>
        </ul>
    </div>
    """
    return HTMLResponse(content=page("My Account", body, username, role))


@app.post("/account/change-password")
async def change_password(request: Request, current_password: str = Form(...), new_password: str = Form(...), confirm_password: str = Form(...)):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)

    if new_password != confirm_password:
        return RedirectResponse("/account?error=New+passwords+do+not+match", status_code=303)

    if len(new_password) < 4:
        return RedirectResponse("/account?error=New+password+must+be+at+least+4+characters", status_code=303)

    if new_password == current_password:
        return RedirectResponse("/account?error=New+password+must+be+different+from+current", status_code=303)

    user_check = verify_login(username, current_password)
    if not user_check:
        return RedirectResponse("/account?error=Current+password+is+incorrect", status_code=303)

    supabase.table("shop_users").update({"password": new_password}).eq("username", username).execute()

    log(username, "password_change", f"User '{username}' changed their password", "warning")

    response = RedirectResponse("/login?error=Password+changed+successfully.+Please+log+in+again", status_code=303)
    response.delete_cookie(COOKIE_NAME)
    return response


def get_cart(request: Request):
    raw = request.cookies.get("cart", "[]")
    try:
        return json.loads(raw)
    except Exception:
        return []


def save_cart(cart):
    response = RedirectResponse("/cart", status_code=303)
    response.set_cookie("cart", json.dumps(cart), max_age=60*60*6)
    return response


# ============ DASHBOARD ============

@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    products = supabase.table("products").select("*").eq("is_active", True).execute().data
    total_products = len(products)
    total_value = sum(float(p.get("quantity_in_stock", 0)) * float(p.get("cost_price", 0)) for p in products)
    low_stock = [p for p in products if float(p.get("quantity_in_stock", 0)) <= float(p.get("reorder_level", 0))]
    low_rows = "".join(f"<tr><td>{p['name']}</td><td class='low'>{p['quantity_in_stock']} {p['unit']}</td></tr>" for p in low_stock)

    customers = supabase.table("customers").select("*").eq("is_active", True).execute().data
    total_owed = sum(float(c.get("balance", 0)) for c in customers)

    chart_section = ""
    charts_script = ""
    if role == "admin":
        all_sales = supabase.table("sales").select("*").execute().data
        all_items = supabase.table("sale_items").select("*").execute().data

        today = datetime.now().date()
        days_30 = [(today - timedelta(days=i)) for i in range(29, -1, -1)]
        daily_labels = [d.strftime('%d/%m') for d in days_30]
        daily_totals = []
        for d in days_30:
            total = 0.0
            for s in all_sales:
                sd = parse_dt(s.get("created_at", ""))
                if sd and sd.date() == d:
                    total += float(s.get("total", 0))
            daily_totals.append(round(total, 2))

        material_rev = {}
        for it in all_items:
            name = it.get("product_name", "Unknown")
            rev = float(it.get("line_total", 0))
            material_rev[name] = material_rev.get(name, 0) + rev
        top_items = sorted(material_rev.items(), key=lambda x: -x[1])[:10]
        top_labels = [t[0] for t in top_items]
        top_values = [round(t[1], 2) for t in top_items]

        payment_totals = {}
        for s in all_sales:
            pm = s.get("payment_method", "Cash")
            payment_totals[pm] = payment_totals.get(pm, 0) + float(s.get("total", 0))
        payment_labels = list(payment_totals.keys())
        payment_values = [round(v, 2) for v in payment_totals.values()]

        month_labels = []
        month_totals = []
        for i in range(5, -1, -1):
            m = today.month - i
            y = today.year
            while m <= 0:
                m += 12
                y -= 1
            month_label = datetime(y, m, 1).strftime('%b %Y')
            month_labels.append(month_label)
            total = 0.0
            for s in all_sales:
                sd = parse_dt(s.get("created_at", ""))
                if sd and sd.year == y and sd.month == m:
                    total += float(s.get("total", 0))
            month_totals.append(round(total, 2))

        chart_section = f"""
        <div class="grid">
            <div class="card">
                <h3>📈 Daily Sales — Last 30 Days</h3>
                <div class="chart-container"><canvas id="dailyChart"></canvas></div>
            </div>
            <div class="card">
                <h3>🥧 Payment Methods (All Time)</h3>
                <div class="chart-container"><canvas id="paymentChart"></canvas></div>
            </div>
        </div>
        <div class="grid">
            <div class="card">
                <h3>📊 Top 10 Materials by Revenue</h3>
                <div class="chart-container"><canvas id="topChart"></canvas></div>
            </div>
            <div class="card">
                <h3>📉 Monthly Sales — Last 6 Months</h3>
                <div class="chart-container"><canvas id="monthlyChart"></canvas></div>
            </div>
        </div>
        """

        charts_script = f"""
        <script src="https://cdn.jsdelivr.net/npm/chart.js"></script>
        <script>
        var dailyLabels = {json.dumps(daily_labels)};
        var dailyData = {json.dumps(daily_totals)};
        var topLabels = {json.dumps(top_labels)};
        var topData = {json.dumps(top_values)};
        var payLabels = {json.dumps(payment_labels)};
        var payData = {json.dumps(payment_values)};
        var monLabels = {json.dumps(month_labels)};
        var monData = {json.dumps(month_totals)};

        new Chart(document.getElementById('dailyChart'), {{
            type: 'line',
            data: {{ labels: dailyLabels, datasets: [{{ label: 'Sales (GHS)', data: dailyData, borderColor: '#1e40af', backgroundColor: 'rgba(30,64,175,0.1)', fill: true, tension: 0.3, pointRadius: 2 }}] }},
            options: {{ responsive: true, maintainAspectRatio: false, plugins: {{ legend: {{ display: false }} }}, scales: {{ y: {{ beginAtZero: true }} }} }}
        }});

        new Chart(document.getElementById('paymentChart'), {{
            type: 'doughnut',
            data: {{ labels: payLabels, datasets: [{{ data: payData, backgroundColor: ['#1e40af','#16a34a','#d97706','#dc2626','#7c3aed','#0891b2'] }}] }},
            options: {{ responsive: true, maintainAspectRatio: false, plugins: {{ legend: {{ position: 'bottom' }} }} }}
        }});

        new Chart(document.getElementById('topChart'), {{
            type: 'bar',
            data: {{ labels: topLabels, datasets: [{{ label: 'Revenue (GHS)', data: topData, backgroundColor: '#1e40af' }}] }},
            options: {{ indexAxis: 'y', responsive: true, maintainAspectRatio: false, plugins: {{ legend: {{ display: false }} }}, scales: {{ x: {{ beginAtZero: true }} }} }}
        }});

        new Chart(document.getElementById('monthlyChart'), {{
            type: 'bar',
            data: {{ labels: monLabels, datasets: [{{ label: 'Sales (GHS)', data: monData, backgroundColor: '#16a34a' }}] }},
            options: {{ responsive: true, maintainAspectRatio: false, plugins: {{ legend: {{ display: false }} }}, scales: {{ y: {{ beginAtZero: true }} }} }}
        }});
        </script>
        """

    admin_actions = ""
    if role == "admin":
        admin_actions = ("<a href='/add' class='btn'>Add Material</a>"
                         "<a href='/reports' class='btn'>📊 Reports</a>"
                         "<a href='/summary' class='btn'>📅 Daily Summary</a>")

    alert_card = ""
    if low_stock:
        alert_card = f"""<div class="card alert-box">
            <h3 style="color:#dc2626;margin-top:0;">🚨 Low Stock Alert</h3>
            <p><strong>{len(low_stock)}</strong> items need restocking. <a href="/alerts" class="btn btn-small btn-danger">View Alerts</a></p>
        </div>"""

    body = f"""
    <h2>Dashboard</h2>
    {alert_card}
    <div class="grid">
        <div class="card stat"><div class="num">{total_products}</div><div class="label">Materials</div></div>
        <div class="card stat"><div class="num">GHS {total_value:,.2f}</div><div class="label">Inventory Value</div></div>
    </div>
    <div class="grid">
        <div class="card stat"><div class="num">{len(customers)}</div><div class="label">Customers</div></div>
        <div class="card stat"><div class="num" style="color:#dc2626;">GHS {total_owed:,.2f}</div><div class="label">Total Owed</div></div>
    </div>
    {chart_section}
    <div class="card">
        <h2>Low Stock ({len(low_stock)})</h2>
        {f"<div class='table-wrap'><table><tr><th>Material</th><th>In Stock</th></tr>{low_rows}</table></div>" if low_stock else "<p>All good!</p>"}
    </div>
    <div class="card">
        {admin_actions}
        <a href="/products" class="btn">View All</a>
        <a href="/sell" class="btn btn-success">New Sale</a>
        <a href="/cart" class="btn">🛒 Cart</a>
        <a href="/customers" class="btn">👥 Customers</a>
        <a href="/suppliers" class="btn">🚚 Suppliers</a>
    </div>
    {charts_script}
    """
    return HTMLResponse(content=page("Dashboard", body, username, role))


# ============ PRODUCTS ============

@app.get("/products", response_class=HTMLResponse)
def products_list(request: Request, search: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    query = supabase.table("products").select("*").eq("is_active", True)
    if search:
        query = query.ilike("name", f"%{search}%")
    products = query.order("name").execute().data

    show_actions = role == "admin"
    rows = ""
    for p in products:
        qty = float(p.get("quantity_in_stock", 0))
        reorder = float(p.get("reorder_level", 0))
        status = "low" if qty <= reorder else "ok"
        actions = ""
        if show_actions:
            actions = (f"<a href='/edit/{p['id']}' class='btn btn-warn btn-small'>✏️</a> "
                       f"<a href='/delete/{p['id']}' class='btn btn-danger btn-small' onclick=\"return confirm('Delete {p['name']}?')\">🗑️</a>")
        rows += f"""<tr>
            <td><strong>{p['name']}</strong><br><small>{p.get('sku','')}</small></td>
            <td>{p.get('unit','')}</td>
            <td class='{status}'>{qty}</td>
            <td>GHS {float(p['selling_price']):,.2f}</td>
            {f"<td>{actions}</td>" if show_actions else ""}
        </tr>"""

    headers = "<tr><th>Material</th><th>Unit</th><th>Stock</th><th>Price</th>"
    if show_actions:
        headers += "<th>Actions</th>"
    headers += "</tr>"
    add_button = "<a href='/add' class='btn'>➕ Add Material</a>" if role == "admin" else ""

    body = f"""
    <h2>All Materials</h2>
    <div class="card">
        {add_button}
        <form method="get" style="display:flex;gap:10px;margin-top:10px;flex-wrap:wrap;">
            <input type="text" name="search" placeholder="Search..." value="{search}" style="flex:1;min-width:200px;">
            <button type="submit">Search</button>
        </form>
    </div>
    <div class="card"><div class="table-wrap"><table>{headers}{rows if rows else "<tr><td colspan='5'>No materials yet.</td></tr>"}</table></div></div>
    """
    return HTMLResponse(content=page("Materials", body, username, role))


@app.get("/add", response_class=HTMLResponse)
def add_form(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add materials")
    cats = supabase.table("categories").select("*").order("name").execute().data
    cat_options = "".join(f"<option value='{c['id']}'>{c['name']}</option>" for c in cats)
    body = f"""
    <h2>Add Material</h2>
    <div class="card">
        <form method="post" action="/add">
            <label>Name</label><input type="text" name="name" required>
            <label>SKU</label><input type="text" name="sku" required>
            <label>Category</label><select name="category_id" required><option value="">--</option>{cat_options}</select>
            <label>Unit</label><input type="text" name="unit" required placeholder="bag">
            <label>Location</label><input type="text" name="location">
            <label>Cost Price</label><input type="number" step="0.01" name="cost_price" required>
            <label>Selling Price</label><input type="number" step="0.01" name="selling_price" required>
            <label>Quantity</label><input type="number" step="0.01" name="quantity_in_stock" required>
            <label>Reorder Level</label><input type="number" step="0.01" name="reorder_level" value="10">
            <button type="submit">Save</button>
            <a href="/products" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Add", body, username, info.get("role")))


@app.post("/add")
async def add_product(request: Request, name: str = Form(...), sku: str = Form(...), category_id: str = Form(...), unit: str = Form(...), location: str = Form(""), cost_price: float = Form(...), selling_price: float = Form(...), quantity_in_stock: float = Form(...), reorder_level: float = Form(10)):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add materials")
    data = {"name": name, "sku": sku, "category_id": int(category_id), "unit": unit, "location": location or None, "cost_price": cost_price, "selling_price": selling_price, "quantity_in_stock": quantity_in_stock, "reorder_level": reorder_level, "is_active": True}
    result = supabase.table("products").insert(data).execute()
    if result.data and quantity_in_stock > 0:
        supabase.table("stock_movements").insert({"product_id": result.data[0]["id"], "movement_type": "IN", "quantity": quantity_in_stock, "unit_cost": cost_price, "note": "Opening stock"}).execute()
    log(username, "material_add", f"Added new material '{name}' (SKU: {sku}, Qty: {quantity_in_stock}, Price: GHS {selling_price})", "info")
    return RedirectResponse("/products", status_code=303)


@app.get("/edit/{product_id}", response_class=HTMLResponse)
def edit_form(request: Request, product_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can edit materials")
    p = supabase.table("products").select("*").eq("id", product_id).single().execute().data
    cats = supabase.table("categories").select("*").order("name").execute().data
    cat_options = "".join(f"<option value='{c['id']}' {'selected' if c['id']==p.get('category_id') else ''}>{c['name']}</option>" for c in cats)
    body = f"""
    <h2>Edit Material</h2>
    <div class="card">
        <form method="post" action="/edit/{product_id}">
            <label>Name</label><input type="text" name="name" value="{p['name']}" required>
            <label>SKU</label><input type="text" name="sku" value="{p.get('sku','')}" required>
            <label>Category</label><select name="category_id" required>{cat_options}</select>
            <label>Unit</label><input type="text" name="unit" value="{p.get('unit','')}" required>
            <label>Location</label><input type="text" name="location" value="{p.get('location','') or ''}">
            <label>Cost Price</label><input type="number" step="0.01" name="cost_price" value="{p['cost_price']}" required>
            <label>Selling Price</label><input type="number" step="0.01" name="selling_price" value="{p['selling_price']}" required>
            <label>Current Stock</label><input type="number" step="0.01" name="quantity_in_stock" value="{p['quantity_in_stock']}" required>
            <label>Reorder Level</label><input type="number" step="0.01" name="reorder_level" value="{p['reorder_level']}">
            <button type="submit">Save Changes</button>
            <a href="/products" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Edit", body, username, info.get("role")))


@app.post("/edit/{product_id}")
async def edit_product(request: Request, product_id: int, name: str = Form(...), sku: str = Form(...), category_id: str = Form(...), unit: str = Form(...), location: str = Form(""), cost_price: float = Form(...), selling_price: float = Form(...), quantity_in_stock: float = Form(...), reorder_level: float = Form(10)):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can edit materials")
    old = supabase.table("products").select("*").eq("id", product_id).single().execute().data
    old_qty = float(old.get("quantity_in_stock", 0))
    old_price = float(old.get("selling_price", 0))
    old_cost = float(old.get("cost_price", 0))

    data = {"name": name, "sku": sku, "category_id": int(category_id), "unit": unit, "location": location or None, "cost_price": cost_price, "selling_price": selling_price, "quantity_in_stock": quantity_in_stock, "reorder_level": reorder_level}
    supabase.table("products").update(data).eq("id", product_id).execute()

    if abs(selling_price - old_price) > 0.001:
        log(username, "price_change",
            f"Changed SELLING price of '{name}' from GHS {old_price:,.2f} to GHS {selling_price:,.2f}",
            "warning")
    if abs(cost_price - old_cost) > 0.001:
        log(username, "cost_change",
            f"Changed COST price of '{name}' from GHS {old_cost:,.2f} to GHS {cost_price:,.2f}",
            "warning")

    if abs(quantity_in_stock - old_qty) > 0.001:
        supabase.table("stock_movements").insert({"product_id": product_id, "movement_type": "ADJUSTMENT", "quantity": quantity_in_stock - old_qty, "note": f"Manual edit by {username}"}).execute()
        log(username, "stock_adjust",
            f"Adjusted stock of '{name}' from {old_qty} to {quantity_in_stock} (change: {quantity_in_stock - old_qty})",
            "warning")

    log(username, "material_edit", f"Edited material '{name}' (SKU: {sku})", "info")
    return RedirectResponse("/products", status_code=303)


@app.get("/delete/{product_id}")
def delete_product(request: Request, product_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can delete materials")
    old = supabase.table("products").select("*").eq("id", product_id).single().execute().data
    supabase.table("products").update({"is_active": False}).eq("id", product_id).execute()
    log(username, "material_delete", f"DELETED material '{old.get('name','')}' (SKU: {old.get('sku','')})", "danger")
    return RedirectResponse("/products", status_code=303)


# ============ PROFIT & LOSS ============

@app.get("/pnl", response_class=HTMLResponse)
def pnl_report(request: Request, start: str = "", end: str = "", preset: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can view P&L")

    now = datetime.now()
    today_str = now.strftime("%Y-%m-%d")

    if preset == "month":
        start = now.replace(day=1).strftime("%Y-%m-%d")
        end = today_str
    elif preset == "lastmonth":
        first_this = now.replace(day=1)
        last_prev = first_this - timedelta(days=1)
        start = last_prev.replace(day=1).strftime("%Y-%m-%d")
        end = last_prev.strftime("%Y-%m-%d")
    elif preset == "quarter":
        q_month = ((now.month - 1) // 3) * 3 + 1
        start = now.replace(month=q_month, day=1).strftime("%Y-%m-%d")
        end = today_str
    elif preset == "year":
        start = now.replace(month=1, day=1).strftime("%Y-%m-%d")
        end = today_str
    elif preset == "all":
        start = "2020-01-01"
        end = today_str

    if not start:
        start = now.replace(day=1).strftime("%Y-%m-%d")
    if not end:
        end = today_str

    def in_range(date_str):
        if not date_str:
            return False
        d = str(date_str)[:10]
        return start <= d <= end

    all_sales = supabase.table("sales").select("*").execute().data
    all_items = supabase.table("sale_items").select("*").execute().data
    all_expenses = supabase.table("expenses").select("*").execute().data

    period_sales = [s for s in all_sales if in_range(s.get("created_at", ""))]
    sale_ids = [s["id"] for s in period_sales]
    period_items = [it for it in all_items if it.get("sale_id") in sale_ids]
    period_expenses = [e for e in all_expenses if in_range(e.get("expense_date", ""))]

    revenue = sum(float(s.get("total", 0)) for s in period_sales)
    cogs = sum(float(it.get("cost_price", 0)) * float(it.get("quantity", 0)) for it in period_items)
    gross_profit = revenue - cogs

    expense_by_cat = {}
    for e in period_expenses:
        cat = e.get("category_name", "Miscellaneous")
        expense_by_cat[cat] = expense_by_cat.get(cat, 0) + float(e.get("amount", 0))
    total_expenses = sum(expense_by_cat.values())
    net_profit = gross_profit - total_expenses

    gross_margin_pct = (gross_profit / revenue * 100) if revenue > 0 else 0
    net_margin_pct = (net_profit / revenue * 100) if revenue > 0 else 0

    expense_rows = ""
    for cat, amt in sorted(expense_by_cat.items(), key=lambda x: -x[1]):
        expense_rows += f"""<tr>
            <td class="label-col">&nbsp;&nbsp;&nbsp;{cat}</td>
            <td class="amount-col">GHS {amt:,.2f}</td>
        </tr>"""
    if not expense_rows:
        expense_rows = '<tr><td class="label-col" colspan="2" style="text-align:center;color:#666;">No expenses in this period</td></tr>'

    net_class = "pnl-profit" if net_profit >= 0 else "pnl-loss"
    net_label = "NET PROFIT" if net_profit >= 0 else "NET LOSS"

    body = f"""
    <div class="card date-bar no-print">
        <h3 style="margin-top:0;">📅 Select Period</h3>
        <form method="get" action="/pnl">
            <div class="field"><label>From</label><input type="date" name="start" value="{start}"></div>
            <div class="field"><label>To</label><input type="date" name="end" value="{end}"></div>
            <button type="submit" class="btn">Apply</button>
        </form>
        <div class="quick-links" style="margin-top:10px;">
            <a href="/pnl?preset=month" class="btn btn-quick">This Month</a>
            <a href="/pnl?preset=lastmonth" class="btn btn-quick">Last Month</a>
            <a href="/pnl?preset=quarter" class="btn btn-quick">This Quarter</a>
            <a href="/pnl?preset=year" class="btn btn-quick">This Year</a>
            <a href="/pnl?preset=all" class="btn btn-quick">All Time</a>
        </div>
    </div>

    <div class="card" style="max-width:850px;margin:0 auto;padding:35px;">
        <div style="text-align:center;border-bottom:3px solid #1e40af;padding-bottom:15px;margin-bottom:25px;">
            <h1 style="color:#1e40af;margin:0;font-size:26px;">🏗️ {SHOP_NAME}</h1>
            <p style="margin:5px 0 0 0;color:#666;">{SHOP_ADDRESS}</p>
            <p style="margin:2px 0 0 0;color:#666;">📞 {SHOP_PHONE}</p>
        </div>

        <h2 style="text-align:center;color:#1e40af;margin-top:0;">PROFIT & LOSS STATEMENT</h2>
        <p style="text-align:center;color:#666;font-size:14px;margin-bottom:25px;">
            Period: <strong>{start}</strong> to <strong>{end}</strong>
        </p>

        <table class="pnl-table">
            <tr class="pnl-section-header">
                <td class="label-col">REVENUE</td>
                <td class="amount-col">AMOUNT (GHS)</td>
            </tr>
            <tr>
                <td class="label-col">&nbsp;&nbsp;&nbsp;Total Sales Revenue ({len(period_sales)} sales)</td>
                <td class="amount-col">{revenue:,.2f}</td>
            </tr>
            <tr class="pnl-subtotal">
                <td class="label-col">TOTAL REVENUE</td>
                <td class="amount-col">{revenue:,.2f}</td>
            </tr>
            <tr class="pnl-section-header">
                <td class="label-col">COST OF GOODS SOLD</td>
                <td class="amount-col"></td>
            </tr>
            <tr>
                <td class="label-col">&nbsp;&nbsp;&nbsp;Cost of materials sold</td>
                <td class="amount-col">({cogs:,.2f})</td>
            </tr>
            <tr class="pnl-subtotal">
                <td class="label-col">GROSS PROFIT</td>
                <td class="amount-col">{gross_profit:,.2f}</td>
            </tr>
            <tr>
                <td class="label-col" style="color:#666;font-size:13px;">&nbsp;&nbsp;&nbsp;Gross Margin: {gross_margin_pct:.1f}%</td>
                <td class="amount-col"></td>
            </tr>
            <tr class="pnl-section-header">
                <td class="label-col">OPERATING EXPENSES</td>
                <td class="amount-col"></td>
            </tr>
            {expense_rows}
            <tr class="pnl-subtotal">
                <td class="label-col">TOTAL EXPENSES</td>
                <td class="amount-col">({total_expenses:,.2f})</td>
            </tr>
            <tr class="pnl-final {net_class}">
                <td class="label-col">{net_label}</td>
                <td class="amount-col">GHS {net_profit:,.2f}</td>
            </tr>
        </table>

        <div style="margin-top:20px;text-align:center;color:#666;font-size:13px;">
            <p style="margin:5px 0;">Net Profit Margin: <strong>{net_margin_pct:.1f}%</strong></p>
        </div>

        <div style="margin-top:30px;padding-top:15px;border-top:1px solid #ddd;text-align:center;color:#666;font-size:12px;">
            <p style="margin:5px 0;">Generated on {now.strftime('%d %B %Y at %H:%M')} by {username}</p>
            <p style="margin:5px 0;">This is a computer-generated statement.</p>
        </div>
    </div>

    <div class="card no-print" style="text-align:center;">
        <button onclick="window.print()" class="btn btn-success">🖨️ Print P&L Statement</button>
        <a href="/export/pnl?start={start}&end={end}" class="btn btn-success">📥 Export to Excel</a>
        <a href="/reports" class="btn">← Back to Reports</a>
    </div>

    <style>
    @media print {{
        .header, .btn, button, .no-print {{ display: none !important; }}
        body {{ background: white; }}
        .card {{ box-shadow: none; padding: 0; margin: 0; }}
        .pnl-table {{ page-break-inside: avoid; }}
        @page {{ margin: 1.5cm; }}
    }}
    </style>
    """
    return HTMLResponse(content=page("Profit & Loss", body, username, info.get("role")))


@app.get("/export/pnl")
def export_pnl(request: Request, start: str = "", end: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export P&L")

    now = datetime.now()
    if not start:
        start = now.replace(day=1).strftime("%Y-%m-%d")
    if not end:
        end = now.strftime("%Y-%m-%d")

    def in_range(date_str):
        if not date_str:
            return False
        d = str(date_str)[:10]
        return start <= d <= end

    all_sales = supabase.table("sales").select("*").execute().data
    all_items = supabase.table("sale_items").select("*").execute().data
    all_expenses = supabase.table("expenses").select("*").execute().data

    period_sales = [s for s in all_sales if in_range(s.get("created_at", ""))]
    sale_ids = [s["id"] for s in period_sales]
    period_items = [it for it in all_items if it.get("sale_id") in sale_ids]
    period_expenses = [e for e in all_expenses if in_range(e.get("expense_date", ""))]

    revenue = sum(float(s.get("total", 0)) for s in period_sales)
    cogs = sum(float(it.get("cost_price", 0)) * float(it.get("quantity", 0)) for it in period_items)
    gross_profit = revenue - cogs

    expense_by_cat = {}
    for e in period_expenses:
        cat = e.get("category_name", "Miscellaneous")
        expense_by_cat[cat] = expense_by_cat.get(cat, 0) + float(e.get("amount", 0))
    total_expenses = sum(expense_by_cat.values())
    net_profit = gross_profit - total_expenses

    wb = Workbook()
    ws = wb.active
    ws.title = "Profit and Loss"

    ws.append([SHOP_NAME])
    ws.append(["PROFIT & LOSS STATEMENT"])
    ws.append([f"Period: {start} to {end}"])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(bold=True, size=12)

    ws.append(["Item", "Amount (GHS)"])
    header_row = ws.max_row
    ws.cell(row=header_row, column=1).font = Font(bold=True, color="FFFFFF")
    ws.cell(row=header_row, column=2).font = Font(bold=True, color="FFFFFF")
    ws.cell(row=header_row, column=1).fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    ws.cell(row=header_row, column=2).fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")

    ws.append(["REVENUE", ""])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([f"  Total Sales ({len(period_sales)} sales)", revenue])
    ws.append(["TOTAL REVENUE", revenue])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)

    ws.append([])
    ws.append(["COST OF GOODS SOLD", ""])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append(["  Cost of materials sold", cogs])
    ws.append(["GROSS PROFIT", gross_profit])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)

    ws.append([])
    ws.append(["OPERATING EXPENSES", ""])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for cat, amt in sorted(expense_by_cat.items(), key=lambda x: -x[1]):
        ws.append([f"  {cat}", amt])
    if not expense_by_cat:
        ws.append(["  (none)", 0])
    ws.append(["TOTAL EXPENSES", total_expenses])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)

    ws.append([])
    ws.append(["NET PROFIT" if net_profit >= 0 else "NET LOSS", net_profit])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True, size=12)
    if net_profit >= 0:
        ws.cell(row=ws.max_row, column=1).fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        ws.cell(row=ws.max_row, column=2).fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    else:
        ws.cell(row=ws.max_row, column=1).fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        ws.cell(row=ws.max_row, column=2).fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    ws.column_dimensions['A'].width = 45
    ws.column_dimensions['B'].width = 18

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_PnL_{start}_to_{end}.xlsx"}
    )


# ============ PROFIT MARGINS ============

@app.get("/margins", response_class=HTMLResponse)
def margins_page(request: Request, filter_type: str = "all", sort_by: str = "margin"):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can view profit margins")

    products = supabase.table("products").select("*").eq("is_active", True).execute().data
    all_items = supabase.table("sale_items").select("*").execute().data

    sold_stats = {}
    for it in all_items:
        pid = it.get("product_id")
        if pid is None:
            continue
        qty = float(it.get("quantity", 0))
        rev = float(it.get("line_total", 0))
        cost = float(it.get("cost_price", 0)) * qty
        if pid not in sold_stats:
            sold_stats[pid] = {"qty": 0, "revenue": 0, "cost": 0}
        sold_stats[pid]["qty"] += qty
        sold_stats[pid]["revenue"] += rev
        sold_stats[pid]["cost"] += cost

    rows_data = []
    for p in products:
        cost = float(p.get("cost_price", 0))
        sell = float(p.get("selling_price", 0))
        profit_per_unit = sell - cost
        margin_pct = (profit_per_unit / sell * 100) if sell > 0 else 0
        stat = sold_stats.get(p["id"], {"qty": 0, "revenue": 0, "cost": 0})
        total_profit = stat["revenue"] - stat["cost"]
        rows_data.append({
            "id": p["id"], "name": p["name"], "sku": p.get("sku", ""), "unit": p.get("unit", ""),
            "cost": cost, "sell": sell, "profit_per_unit": profit_per_unit, "margin_pct": margin_pct,
            "qty_sold": stat["qty"], "revenue": stat["revenue"], "total_profit": total_profit,
        })

    if filter_type == "profit":
        rows_data = [r for r in rows_data if r["profit_per_unit"] > 0]
    elif filter_type == "loss":
        rows_data = [r for r in rows_data if r["profit_per_unit"] <= 0]
    elif filter_type == "lowmargin":
        rows_data = [r for r in rows_data if 0 < r["margin_pct"] < 10]
    elif filter_type == "highmargin":
        rows_data = [r for r in rows_data if r["margin_pct"] >= 30]

    if sort_by == "margin":
        rows_data.sort(key=lambda x: x["margin_pct"], reverse=True)
    elif sort_by == "profit":
        rows_data.sort(key=lambda x: x["total_profit"], reverse=True)
    elif sort_by == "qty":
        rows_data.sort(key=lambda x: x["qty_sold"], reverse=True)
    elif sort_by == "name":
        rows_data.sort(key=lambda x: x["name"].lower())

    total_materials = len(products)
    avg_margin = (sum(r["margin_pct"] for r in rows_data) / len(rows_data)) if rows_data else 0
    profitable_count = sum(1 for r in rows_data if r["profit_per_unit"] > 0)
    loss_count = sum(1 for r in rows_data if r["profit_per_unit"] <= 0)

    best = max(rows_data, key=lambda x: x["margin_pct"], default=None)
    worst = min(rows_data, key=lambda x: x["margin_pct"], default=None)

    rows = ""
    for r in rows_data:
        if r["margin_pct"] >= 30:
            margin_class = "good-margin"
        elif r["margin_pct"] >= 10:
            margin_class = "ok-margin"
        else:
            margin_class = "bad-margin"
        rows += f"""<tr>
            <td><strong>{r['name']}</strong><br><small>{r['sku']}</small></td>
            <td>{r['unit']}</td>
            <td>GHS {r['cost']:,.2f}</td>
            <td>GHS {r['sell']:,.2f}</td>
            <td class='{margin_class}'>GHS {r['profit_per_unit']:,.2f}</td>
            <td class='{margin_class}'>{r['margin_pct']:.1f}%</td>
            <td>{r['qty_sold']:.0f}</td>
            <td>GHS {r['total_profit']:,.2f}</td>
        </tr>"""

    best_html = f"<p><strong>🏆 Best Margin:</strong> {best['name']} — <span class='good-margin'>{best['margin_pct']:.1f}%</span></p>" if best else ""
    worst_html = f"<p><strong>⚠️ Lowest Margin:</strong> {worst['name']} — <span class='bad-margin'>{worst['margin_pct']:.1f}%</span></p>" if worst else ""

    body = f"""
    <h2>📈 Profit Margins per Material</h2>
    <div class="grid">
        <div class="card stat"><div class="num">{total_materials}</div><div class="label">Total Materials</div></div>
        <div class="card stat"><div class="num">{avg_margin:.1f}%</div><div class="label">Average Margin</div></div>
    </div>
    <div class="grid">
        <div class="card stat"><div class="num" style="color:#16a34a;">{profitable_count}</div><div class="label">Profitable Items</div></div>
        <div class="card stat"><div class="num" style="color:#dc2626;">{loss_count}</div><div class="label">Loss / No Profit</div></div>
    </div>
    <div class="card">{best_html}{worst_html}</div>
    <div class="card">
        <h3>🎛️ Filters & Sort</h3>
        <form method="get" action="/margins" style="display:flex;gap:10px;flex-wrap:wrap;align-items:flex-end;">
            <div style="flex:1;min-width:180px;">
                <label>Show</label>
                <select name="filter_type">
                    <option value="all" {'selected' if filter_type=='all' else ''}>All Materials</option>
                    <option value="profit" {'selected' if filter_type=='profit' else ''}>Only Profitable</option>
                    <option value="loss" {'selected' if filter_type=='loss' else ''}>Only Loss/No Profit</option>
                    <option value="lowmargin" {'selected' if filter_type=='lowmargin' else ''}>Low Margin (below 10%)</option>
                    <option value="highmargin" {'selected' if filter_type=='highmargin' else ''}>High Margin (30%+)</option>
                </select>
            </div>
            <div style="flex:1;min-width:180px;">
                <label>Sort By</label>
                <select name="sort_by">
                    <option value="margin" {'selected' if sort_by=='margin' else ''}>Margin % (high → low)</option>
                    <option value="profit" {'selected' if sort_by=='profit' else ''}>Total Profit (high → low)</option>
                    <option value="qty" {'selected' if sort_by=='qty' else ''}>Qty Sold (high → low)</option>
                    <option value="name" {'selected' if sort_by=='name' else ''}>Name (A → Z)</option>
                </select>
            </div>
            <button type="submit" class="btn">Apply</button>
            <a href="/export/margins" class="btn btn-success">📥 Export Excel</a>
        </form>
    </div>
    <div class="card">
        <h3>📊 Materials List ({len(rows_data)})</h3>
        <div class="table-wrap">
        <table>
            <tr><th>Material</th><th>Unit</th><th>Cost</th><th>Price</th><th>Profit/Unit</th><th>Margin %</th><th>Qty Sold</th><th>Total Profit</th></tr>
            {rows if rows else "<tr><td colspan='8'>No materials match this filter.</td></tr>"}
        </table>
        </div>
        <p style="margin-top:15px;color:#666;font-size:13px;">🟢 30%+ margin · 🟡 10–29% margin · 🔴 Below 10% or loss</p>
    </div>
    """
    return HTMLResponse(content=page("Margins", body, username, info.get("role")))


@app.get("/export/margins")
def export_margins(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export data")

    products = supabase.table("products").select("*").eq("is_active", True).execute().data
    all_items = supabase.table("sale_items").select("*").execute().data

    sold_stats = {}
    for it in all_items:
        pid = it.get("product_id")
        if pid is None:
            continue
        qty = float(it.get("quantity", 0))
        rev = float(it.get("line_total", 0))
        cost = float(it.get("cost_price", 0)) * qty
        if pid not in sold_stats:
            sold_stats[pid] = {"qty": 0, "revenue": 0, "cost": 0}
        sold_stats[pid]["qty"] += qty
        sold_stats[pid]["revenue"] += rev
        sold_stats[pid]["cost"] += cost

    wb = Workbook()
    ws = wb.active
    ws.title = "Profit Margins"
    headers = ["Material", "SKU", "Unit", "Cost", "Price", "Profit/Unit", "Margin %", "Qty Sold", "Revenue", "Total Profit"]
    style_header(ws, headers)

    for p in products:
        cost = float(p.get("cost_price", 0))
        sell = float(p.get("selling_price", 0))
        profit_per_unit = sell - cost
        margin_pct = (profit_per_unit / sell * 100) if sell > 0 else 0
        stat = sold_stats.get(p["id"], {"qty": 0, "revenue": 0, "cost": 0})
        total_profit = stat["revenue"] - stat["cost"]
        ws.append([p.get("name", ""), p.get("sku", ""), p.get("unit", ""), cost, sell, profit_per_unit, round(margin_pct, 1), stat["qty"], stat["revenue"], total_profit])

    widths = [30, 15, 10, 12, 12, 14, 12, 12, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_margins_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"}
    )


# ============ CUSTOMERS ============

@app.get("/customers", response_class=HTMLResponse)
def customers_list(request: Request, search: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    query = supabase.table("customers").select("*").eq("is_active", True)
    if search:
        query = query.ilike("name", f"%{search}%")
    customers = query.order("name").execute().data

    rows = ""
    for c in customers:
        balance = float(c.get("balance", 0))
        balance_class = "owed" if balance > 0 else "clear"
        rows += f"""<tr>
            <td><strong>{c['name']}</strong><br><small>{c.get('phone','')}</small></td>
            <td>{c.get('address','') or ''}</td>
            <td class='{balance_class}'>GHS {balance:,.2f}</td>
            <td>
                <a href='/customers/view/{c['id']}' class='btn btn-small'>👁️ View</a>
                <a href='/customers/pay/{c['id']}' class='btn btn-success btn-small'>💰 Pay</a>
            </td>
        </tr>"""

    total_owed = sum(float(c.get("balance", 0)) for c in customers)

    body = f"""
    <h2>👥 Customers</h2>
    <div class="card"><h3>Total Owed: <span style="color:#dc2626;">GHS {total_owed:,.2f}</span></h3></div>
    <div class="card"><a href="/customers/add" class="btn btn-success">➕ Add Customer</a></div>
    <div class="card">
        <form method="get" style="display:flex;gap:10px;flex-wrap:wrap;">
            <input type="text" name="search" placeholder="Search by name..." value="{search}" style="flex:1;min-width:200px;">
            <button type="submit">Search</button>
        </form>
    </div>
    <div class="card">
        <div class="table-wrap">
        <table>
            <tr><th>Customer</th><th>Address</th><th>Balance Owed</th><th>Actions</th></tr>
            {rows if rows else "<tr><td colspan='4'>No customers yet.</td></tr>"}
        </table>
        </div>
    </div>
    """
    return HTMLResponse(content=page("Customers", body, username, role))


@app.get("/customers/add", response_class=HTMLResponse)
def customer_add_form(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"
    body = """
    <h2>➕ Add New Customer</h2>
    <div class="card">
        <form method="post" action="/customers/add">
            <label>Full Name</label>
            <input type="text" name="name" required>
            <label>Phone Number</label>
            <input type="text" name="phone" placeholder="e.g. 0244123456">
            <label>Address / Location</label>
            <input type="text" name="address" placeholder="e.g. Menzezor">
            <label>Notes (optional)</label>
            <input type="text" name="notes">
            <button type="submit" class="btn btn-success">Create Customer</button>
            <a href="/customers" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Add Customer", body, username, role))


@app.post("/customers/add")
async def customer_add(request: Request, name: str = Form(...), phone: str = Form(""), address: str = Form(""), notes: str = Form("")):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    supabase.table("customers").insert({"name": name, "phone": phone or None, "address": address or None, "notes": notes or None, "balance": 0, "is_active": True}).execute()
    log(username, "customer_add", f"Added new customer '{name}' (Phone: {phone or '—'})", "info")
    return RedirectResponse("/customers", status_code=303)


@app.get("/customers/view/{customer_id}", response_class=HTMLResponse)
def customer_view(request: Request, customer_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    c = supabase.table("customers").select("*").eq("id", customer_id).single().execute().data
    balance = float(c.get("balance", 0))

    sales = supabase.table("sales").select("*").eq("customer_id", customer_id).order("created_at", desc=True).execute().data
    sales_rows = "".join(f"<tr><td>{s.get('created_at','')[:16]}</td><td>{s.get('invoice_no','')}</td><td>GHS {float(s.get('total',0)):,.2f}</td><td>GHS {float(s.get('amount_paid_now',0)):,.2f}</td><td>GHS {float(s.get('amount_on_credit',0)):,.2f}</td></tr>" for s in sales)

    payments = supabase.table("customer_payments").select("*").eq("customer_id", customer_id).order("created_at", desc=True).execute().data
    pay_rows = "".join(f"<tr><td>{p.get('created_at','')[:16]}</td><td>GHS {float(p.get('amount',0)):,.2f}</td><td>{p.get('payment_method','')}</td><td>{p.get('note','') or ''}</td></tr>" for p in payments)

    body = f"""
    <h2>👤 {c['name']}</h2>
    <div class="card">
        <p><strong>Phone:</strong> {c.get('phone','') or '—'}</p>
        <p><strong>Address:</strong> {c.get('address','') or '—'}</p>
        <p><strong>Notes:</strong> {c.get('notes','') or '—'}</p>
        <h3>Current Balance: <span class="{'owed' if balance > 0 else 'clear'}">GHS {balance:,.2f}</span></h3>
        <a href="/customers/pay/{customer_id}" class="btn btn-success">💰 Record Payment</a>
        <a href="/customers/statement/{customer_id}" class="btn">📄 Print Statement</a>
    </div>
    <div class="card">
        <h3>📋 Purchase History</h3>
        <div class="table-wrap"><table>
            <tr><th>Date</th><th>Invoice</th><th>Total</th><th>Paid</th><th>Credit</th></tr>
            {sales_rows if sales_rows else "<tr><td colspan='5'>No purchases yet.</td></tr>"}
        </table></div>
    </div>
    <div class="card">
        <h3>💵 Payment History</h3>
        <div class="table-wrap"><table>
            <tr><th>Date</th><th>Amount</th><th>Method</th><th>Note</th></tr>
            {pay_rows if pay_rows else "<tr><td colspan='4'>No payments yet.</td></tr>"}
        </table></div>
    </div>
    <div class="card"><a href="/customers" class="btn">← Back to Customers</a></div>
    """
    return HTMLResponse(content=page("Customer", body, username, role))


@app.get("/customers/statement/{customer_id}", response_class=HTMLResponse)
def customer_statement(request: Request, customer_id: int, start: str = "", end: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    now = datetime.now()
    if not start:
        start = now.replace(day=1).strftime("%Y-%m-%d")
    if not end:
        end = now.strftime("%Y-%m-%d")

    c = supabase.table("customers").select("*").eq("id", customer_id).single().execute().data
    balance = float(c.get("balance", 0))

    all_sales = supabase.table("sales").select("*").eq("customer_id", customer_id).execute().data
    all_payments = supabase.table("customer_payments").select("*").eq("customer_id", customer_id).execute().data

    def in_range(date_str):
        if not date_str:
            return False
        d = str(date_str)[:10]
        return start <= d <= end

    period_sales = [s for s in all_sales if in_range(s.get("created_at", ""))]
    period_payments = [p for p in all_payments if in_range(p.get("created_at", ""))]

    transactions = []
    for s in period_sales:
        transactions.append({"date": str(s.get("created_at", ""))[:10], "type": "Sale", "ref": s.get("invoice_no", ""), "debit": float(s.get("total", 0)), "credit": float(s.get("amount_paid_now", 0)) if float(s.get("amount_on_credit", 0)) > 0 else 0})
    for p in period_payments:
        transactions.append({"date": str(p.get("created_at", ""))[:10], "type": "Payment", "ref": p.get("payment_method", ""), "debit": 0, "credit": float(p.get("amount", 0))})

    transactions.sort(key=lambda x: x["date"])
    period_debit = sum(t["debit"] for t in transactions)
    period_credit = sum(t["credit"] for t in transactions)

    opening_balance = 0.0
    for s in all_sales:
        d = str(s.get("created_at", ""))[:10]
        if d < start:
            opening_balance += float(s.get("amount_on_credit", 0))
    for p in all_payments:
        d = str(p.get("created_at", ""))[:10]
        if d < start:
            opening_balance -= float(p.get("amount", 0))

    running = opening_balance
    tx_rows = ""
    if not transactions:
        tx_rows = "<tr><td colspan='6' style='text-align:center;color:#666;'>No transactions in this period</td></tr>"
    else:
        for t in transactions:
            running += t["debit"] - t["credit"]
            debit_str = f"GHS {t['debit']:,.2f}" if t["debit"] > 0 else "—"
            credit_str = f"GHS {t['credit']:,.2f}" if t["credit"] > 0 else "—"
            tx_rows += f"<tr><td>{t['date']}</td><td>{t['type']}</td><td>{t['ref']}</td><td style='text-align:right;'>{debit_str}</td><td style='text-align:right;'>{credit_str}</td><td style='text-align:right;'>GHS {running:,.2f}</td></tr>"

    body = f"""
    <div style="max-width:800px;margin:0 auto;">
        <div class="card" id="statement" style="padding:30px;">
            <div style="text-align:center;border-bottom:3px solid #1e40af;padding-bottom:15px;margin-bottom:20px;">
                <h1 style="color:#1e40af;margin:0;font-size:26px;">🏗️ {SHOP_NAME}</h1>
                <p style="margin:5px 0 0 0;color:#666;">{SHOP_ADDRESS}</p>
                <p style="margin:2px 0 0 0;color:#666;">📞 {SHOP_PHONE}</p>
            </div>
            <h2 style="text-align:center;color:#1e40af;margin-top:0;">CUSTOMER STATEMENT</h2>
            <p style="text-align:center;color:#666;font-size:14px;">Period: <strong>{start}</strong> to <strong>{end}</strong></p>
            <div style="display:grid;grid-template-columns:1fr 1fr;gap:20px;margin-bottom:20px;">
                <div>
                    <p style="margin:5px 0;font-size:14px;"><strong>Customer:</strong> {c['name']}</p>
                    <p style="margin:5px 0;font-size:14px;"><strong>Phone:</strong> {c.get('phone','') or '—'}</p>
                    <p style="margin:5px 0;font-size:14px;"><strong>Address:</strong> {c.get('address','') or '—'}</p>
                </div>
                <div style="text-align:right;">
                    <p style="margin:5px 0;font-size:14px;"><strong>Statement Date:</strong> {now.strftime('%d %B %Y')}</p>
                    <p style="margin:5px 0;font-size:14px;"><strong>Prepared By:</strong> {username}</p>
                </div>
            </div>
            <div class="table-wrap">
            <table style="min-width:100%;font-size:14px;">
                <thead>
                    <tr style="background:#1e40af;color:white;">
                        <th style="padding:10px;">Date</th>
                        <th style="padding:10px;">Type</th>
                        <th style="padding:10px;">Ref</th>
                        <th style="padding:10px;text-align:right;">Debit (GHS)</th>
                        <th style="padding:10px;text-align:right;">Credit (GHS)</th>
                        <th style="padding:10px;text-align:right;">Balance</th>
                    </tr>
                </thead>
                <tbody>
                    <tr style="background:#f3f4f6;">
                        <td style="padding:10px;">{start}</td>
                        <td colspan="4" style="padding:10px;"><strong>Opening Balance</strong></td>
                        <td style="padding:10px;text-align:right;"><strong>GHS {opening_balance:,.2f}</strong></td>
                    </tr>
                    {tx_rows}
                    <tr style="background:#fef3c7;font-weight:bold;">
                        <td colspan="3" style="padding:12px;text-align:right;">TOTALS:</td>
                        <td style="padding:12px;text-align:right;">GHS {period_debit:,.2f}</td>
                        <td style="padding:12px;text-align:right;">GHS {period_credit:,.2f}</td>
                        <td style="padding:12px;text-align:right;">GHS {running:,.2f}</td>
                    </tr>
                </tbody>
            </table>
            </div>
            <div style="margin-top:25px;padding:20px;background:{'#fee2e2' if balance > 0 else '#dcfce7'};border-radius:8px;text-align:center;">
                <p style="margin:0;font-size:14px;color:#666;">Current Balance Owed</p>
                <p style="margin:10px 0 0 0;font-size:32px;font-weight:bold;color:{'#dc2626' if balance > 0 else '#16a34a'};">GHS {balance:,.2f}</p>
            </div>
            <div style="margin-top:25px;padding-top:15px;border-top:1px solid #ddd;text-align:center;color:#666;font-size:13px;">
                <p style="margin:5px 0;">This is a computer-generated statement. Please retain for your records.</p>
                <p style="margin:5px 0;">For questions, call <strong>{SHOP_PHONE}</strong></p>
            </div>
        </div>
        <div style="text-align:center;margin-top:15px;" class="no-print">
            <button onclick="window.print()" class="btn btn-success">🖨️ Print Statement</button>
            <a href="/customers/view/{customer_id}" class="btn">← Back to Customer</a>
        </div>
        <div class="card no-print" style="margin-top:15px;">
            <h3>📅 Change Period</h3>
            <form method="get" action="/customers/statement/{customer_id}" style="display:flex;gap:10px;flex-wrap:wrap;align-items:flex-end;">
                <div style="flex:1;min-width:150px;"><label>Start Date</label><input type="date" name="start" value="{start}"></div>
                <div style="flex:1;min-width:150px;"><label>End Date</label><input type="date" name="end" value="{end}"></div>
                <button type="submit" class="btn">Update Statement</button>
            </form>
        </div>
    </div>
    """
    return HTMLResponse(content=page("Statement", body, username, role))


@app.get("/customers/pay/{customer_id}", response_class=HTMLResponse)
def customer_pay_form(request: Request, customer_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"
    c = supabase.table("customers").select("*").eq("id", customer_id).single().execute().data
    balance = float(c.get("balance", 0))
    body = f"""
    <h2>💰 Record Payment from {c['name']}</h2>
    <div class="card">
        <p>Current balance: <strong class="owed">GHS {balance:,.2f}</strong></p>
        <form method="post" action="/customers/pay/{customer_id}">
            <label>Amount Paid (GHS)</label>
            <input type="number" step="0.01" name="amount" required min="0.01" max="{balance}">
            <label>Payment Method</label>
            <select name="payment_method">
                <option value="Cash">💵 Cash</option>
                <option value="Mobile Money">📱 Mobile Money</option>
                <option value="Bank Transfer">🏦 Bank Transfer</option>
                <option value="Card">💳 Card</option>
            </select>
            <label>Note (optional)</label>
            <input type="text" name="note">
            <button type="submit" class="btn btn-success">Record Payment</button>
            <a href="/customers/view/{customer_id}" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Record Payment", body, username, role))


@app.post("/customers/pay/{customer_id}")
async def customer_pay(request: Request, customer_id: int, amount: float = Form(...), payment_method: str = Form("Cash"), note: str = Form("")):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)

    c = supabase.table("customers").select("*").eq("id", customer_id).single().execute().data
    balance = float(c.get("balance", 0))
    if amount > balance:
        amount = balance
    new_balance = balance - amount

    supabase.table("customers").update({"balance": new_balance}).eq("id", customer_id).execute()
    supabase.table("customer_payments").insert({"customer_id": customer_id, "amount": amount, "payment_method": payment_method, "note": note or None, "recorded_by": username}).execute()

    log(username, "customer_payment", f"Recorded payment of GHS {amount:,.2f} from customer '{c.get('name','')}' via {payment_method}", "info")
    return RedirectResponse(f"/customers/view/{customer_id}", status_code=303)


# ============ SUPPLIERS ============

@app.get("/suppliers", response_class=HTMLResponse)
def suppliers_list(request: Request, search: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    query = supabase.table("suppliers").select("*")
    try:
        query = query.eq("is_active", True)
    except Exception:
        pass
    if search:
        query = query.ilike("name", f"%{search}%")
    suppliers = query.order("name").execute().data

    purchase_data = supabase.table("purchase_orders").select("*").execute().data
    totals = {}
    counts = {}
    for po in purchase_data:
        sid = po.get("supplier_id")
        totals[sid] = totals.get(sid, 0) + float(po.get("total", 0))
        counts[sid] = counts.get(sid, 0) + 1

    rows = ""
    for s in suppliers:
        spent = totals.get(s["id"], 0)
        num_orders = counts.get(s["id"], 0)
        rows += f"""<tr>
            <td><strong>{s['name']}</strong><br><small>{s.get('phone','')}</small></td>
            <td>{s.get('address','') or ''}</td>
            <td>{num_orders} orders</td>
            <td>GHS {spent:,.2f}</td>
            <td>
                <a href='/suppliers/view/{s['id']}' class='btn btn-small'>👁️ View</a>
                <a href='/purchases/new?supplier_id={s['id']}' class='btn btn-success btn-small'>📦 Restock</a>
            </td>
        </tr>"""

    body = f"""
    <h2>🚚 Suppliers</h2>
    <div class="card">
        <a href="/suppliers/add" class="btn btn-success">➕ Add Supplier</a>
        <a href="/purchases" class="btn">📦 All Purchases</a>
    </div>
    <div class="card">
        <form method="get" style="display:flex;gap:10px;flex-wrap:wrap;">
            <input type="text" name="search" placeholder="Search suppliers..." value="{search}" style="flex:1;min-width:200px;">
            <button type="submit">Search</button>
        </form>
    </div>
    <div class="card">
        <div class="table-wrap"><table>
            <tr><th>Supplier</th><th>Address</th><th>Orders</th><th>Total Spent</th><th>Actions</th></tr>
            {rows if rows else "<tr><td colspan='5'>No suppliers yet. Add one to start tracking purchases.</td></tr>"}
        </table></div>
    </div>
    """
    return HTMLResponse(content=page("Suppliers", body, username, role))


@app.get("/suppliers/add", response_class=HTMLResponse)
def supplier_add_form(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add suppliers")
    body = """
    <h2>➕ Add New Supplier</h2>
    <div class="card">
        <form method="post" action="/suppliers/add">
            <label>Supplier Name</label>
            <input type="text" name="name" required placeholder="e.g. Dangote Cement Depot">
            <label>Phone Number</label>
            <input type="text" name="phone" placeholder="e.g. 0244123456">
            <label>Email (optional)</label>
            <input type="text" name="email">
            <label>Address</label>
            <input type="text" name="address" placeholder="e.g. Accra">
            <label>Notes (optional)</label>
            <input type="text" name="notes">
            <button type="submit" class="btn btn-success">Create Supplier</button>
            <a href="/suppliers" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Add Supplier", body, username, info.get("role")))


@app.post("/suppliers/add")
async def supplier_add(request: Request, name: str = Form(...), phone: str = Form(""), email: str = Form(""), address: str = Form(""), notes: str = Form("")):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add suppliers")
    supabase.table("suppliers").insert({"name": name, "phone": phone or None, "email": email or None, "address": address or None, "notes": notes or None, "is_active": True}).execute()
    log(username, "supplier_add", f"Added new supplier '{name}' (Phone: {phone or '—'})", "info")
    return RedirectResponse("/suppliers", status_code=303)


@app.get("/suppliers/view/{supplier_id}", response_class=HTMLResponse)
def supplier_view(request: Request, supplier_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    s = supabase.table("suppliers").select("*").eq("id", supplier_id).single().execute().data
    purchases = supabase.table("purchase_orders").select("*").eq("supplier_id", supplier_id).order("created_at", desc=True).execute().data

    rows = ""
    total_spent = 0
    for po in purchases:
        total_spent += float(po.get("total", 0))
        rows += f"<tr><td>{po.get('created_at','')[:16]}</td><td>{po.get('po_number','')}</td><td>GHS {float(po.get('total',0)):,.2f}</td><td>{po.get('payment_method','')}</td><td><a href='/purchases/view/{po['id']}' class='btn btn-small'>View</a></td></tr>"

    body = f"""
    <h2>🚚 {s['name']}</h2>
    <div class="card">
        <p><strong>Phone:</strong> {s.get('phone','') or '—'}</p>
        <p><strong>Email:</strong> {s.get('email','') or '—'}</p>
        <p><strong>Address:</strong> {s.get('address','') or '—'}</p>
        <p><strong>Notes:</strong> {s.get('notes','') or '—'}</p>
        <h3>Total Purchased: <span style="color:#dc2626;">GHS {total_spent:,.2f}</span></h3>
        <a href="/purchases/new?supplier_id={supplier_id}" class="btn btn-success">📦 Record New Purchase</a>
    </div>
    <div class="card">
        <h3>📋 Purchase History ({len(purchases)})</h3>
        <div class="table-wrap"><table>
            <tr><th>Date</th><th>PO Number</th><th>Total</th><th>Payment</th><th></th></tr>
            {rows if rows else "<tr><td colspan='5'>No purchases yet.</td></tr>"}
        </table></div>
    </div>
    <div class="card"><a href="/suppliers" class="btn">← Back to Suppliers</a></div>
    """
    return HTMLResponse(content=page("Supplier", body, username, role))


# ============ PURCHASES ============

@app.get("/purchases", response_class=HTMLResponse)
def purchases_list(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    purchases = supabase.table("purchase_orders").select("*").order("created_at", desc=True).limit(100).execute().data
    rows = ""
    total_all = 0
    for po in purchases:
        total_all += float(po.get("total", 0))
        rows += f"<tr><td>{po.get('created_at','')[:16]}</td><td>{po.get('po_number','')}</td><td>{po.get('supplier_name','')}</td><td>GHS {float(po.get('total',0)):,.2f}</td><td>{po.get('payment_method','')}</td><td><a href='/purchases/view/{po['id']}' class='btn btn-small'>View</a></td></tr>"

    body = f"""
    <h2>📦 Purchase Orders</h2>
    <div class="card">
        <a href="/purchases/new" class="btn btn-success">➕ Record New Purchase</a>
        <a href="/suppliers" class="btn">🚚 Suppliers</a>
    </div>
    <div class="card"><h3>Total Spent: <span style="color:#dc2626;">GHS {total_all:,.2f}</span> ({len(purchases)} orders)</h3></div>
    <div class="card">
        <div class="table-wrap"><table>
            <tr><th>Date</th><th>PO Number</th><th>Supplier</th><th>Total</th><th>Payment</th><th></th></tr>
            {rows if rows else "<tr><td colspan='6'>No purchases yet.</td></tr>"}
        </table></div>
    </div>
    """
    return HTMLResponse(content=page("Purchases", body, username, role))


@app.get("/purchases/new", response_class=HTMLResponse)
def purchase_new_form(request: Request, supplier_id: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can record purchases")

    suppliers = supabase.table("suppliers").select("*").order("name").execute().data
    sup_options = "".join(f"<option value='{s['id']}' {'selected' if str(s['id'])==str(supplier_id) else ''}>{s['name']}</option>" for s in suppliers)

    products = supabase.table("products").select("*").eq("is_active", True).order("name").execute().data
    prod_options = "".join(f"<option value='{p['id']}' data-cost='{p.get('cost_price',0)}'>{p['name']} ({p.get('unit','')})</option>" for p in products)

    body = f"""
    <h2>📦 Record New Purchase</h2>
    <div class="card">
        <form method="post" action="/purchases/new">
            <label>Supplier</label>
            <select name="supplier_id" required><option value="">-- Select supplier --</option>{sup_options}</select>
            <h3>Items Purchased</h3>
            <p style="color:#666;font-size:14px;">Add each item you bought. Stock will be updated automatically.</p>
            <div id="items-container">
                <div class="item-row" style="border:1px solid #ddd;padding:15px;border-radius:8px;margin-bottom:10px;">
                    <label>Material</label>
                    <select name="product_id[]" required><option value="">-- Select material --</option>{prod_options}</select>
                    <label>Quantity</label>
                    <input type="number" step="0.01" name="quantity[]" required min="0.01">
                    <label>Unit Cost (GHS)</label>
                    <input type="number" step="0.01" name="unit_cost[]" required min="0.01">
                </div>
            </div>
            <button type="button" onclick="addItem()" class="btn">➕ Add Another Item</button>
            <h3 style="margin-top:20px;">Payment</h3>
            <label>Payment Method</label>
            <select name="payment_method">
                <option value="Cash">💵 Cash</option>
                <option value="Mobile Money">📱 Mobile Money</option>
                <option value="Bank Transfer">🏦 Bank Transfer</option>
                <option value="Card">💳 Card</option>
                <option value="Credit">📝 Credit (I Owe Supplier)</option>
            </select>
            <label>Amount Paid Now</label>
            <input type="number" step="0.01" name="amount_paid" value="0" min="0">
            <label>Note (optional)</label>
            <input type="text" name="note">
            <button type="submit" class="btn btn-success">✅ Save Purchase</button>
            <a href="/purchases" class="btn">Cancel</a>
        </form>
    </div>
    <script>
    function addItem() {{
        var container = document.getElementById('items-container');
        var firstRow = container.querySelector('.item-row');
        var newRow = firstRow.cloneNode(true);
        newRow.querySelectorAll('input').forEach(function(inp) {{ inp.value = ''; }});
        newRow.querySelectorAll('select').forEach(function(sel) {{ sel.selectedIndex = 0; }});
        container.appendChild(newRow);
    }}
    </script>
    """
    return HTMLResponse(content=page("New Purchase", body, username, info.get("role")))


@app.post("/purchases/new")
async def purchase_new(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can record purchases")

    form = await request.form()
    supplier_id = form.get("supplier_id")
    if not supplier_id:
        raise HTTPException(400, "Supplier is required")

    supplier = supabase.table("suppliers").select("*").eq("id", int(supplier_id)).single().execute().data
    supplier_name = supplier["name"] if supplier else "Unknown"

    product_ids = form.getlist("product_id[]")
    quantities = form.getlist("quantity[]")
    unit_costs = form.getlist("unit_cost[]")

    if not product_ids or len(product_ids) != len(quantities):
        raise HTTPException(400, "Invalid items")

    payment_method = form.get("payment_method", "Cash")
    amount_paid = float(form.get("amount_paid", 0) or 0)
    note = form.get("note", "")

    items = []
    total = 0.0
    for i in range(len(product_ids)):
        if not product_ids[i]:
            continue
        pid = int(product_ids[i])
        qty = float(quantities[i])
        cost = float(unit_costs[i])
        line = qty * cost
        total += line
        product = supabase.table("products").select("*").eq("id", pid).single().execute().data
        items.append({"product_id": pid, "product_name": product["name"], "quantity": qty, "unit_cost": cost, "line_total": line})

    if not items:
        raise HTTPException(400, "At least one item required")

    credit_amount = 0.0
    if payment_method == "Credit":
        credit_amount = total - amount_paid
        if credit_amount < 0:
            credit_amount = 0

    po_number = f"PO-{datetime.now().strftime('%Y%m%d%H%M%S')}"
    po_data = {"po_number": po_number, "supplier_id": int(supplier_id), "supplier_name": supplier_name, "total": total, "amount_paid": amount_paid if payment_method == "Credit" else total, "amount_owed": credit_amount, "payment_method": payment_method, "status": "received", "note": note or None, "recorded_by": username}
    result = supabase.table("purchase_orders").insert(po_data).execute()
    purchase_id = result.data[0]["id"]

    for it in items:
        supabase.table("purchase_order_items").insert({"purchase_id": purchase_id, "product_id": it["product_id"], "product_name": it["product_name"], "quantity": it["quantity"], "unit_cost": it["unit_cost"], "line_total": it["line_total"]}).execute()
        p = supabase.table("products").select("*").eq("id", it["product_id"]).single().execute().data
        new_qty = float(p.get("quantity_in_stock", 0)) + it["quantity"]
        supabase.table("products").update({"quantity_in_stock": new_qty, "cost_price": it["unit_cost"]}).eq("id", it["product_id"]).execute()
        supabase.table("stock_movements").insert({"product_id": it["product_id"], "movement_type": "IN", "quantity": it["quantity"], "unit_cost": it["unit_cost"], "reference": po_number, "note": f"Purchase from {supplier_name} — {it['product_name']} x {it['quantity']}"}).execute()

    log(username, "purchase_add", f"Recorded purchase from '{supplier_name}' - Total GHS {total:,.2f} ({len(items)} items)", "info")
    return RedirectResponse(f"/purchases/view/{purchase_id}", status_code=303)


@app.get("/purchases/view/{purchase_id}", response_class=HTMLResponse)
def purchase_view(request: Request, purchase_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"

    po = supabase.table("purchase_orders").select("*").eq("id", purchase_id).single().execute().data
    items = supabase.table("purchase_order_items").select("*").eq("purchase_id", purchase_id).execute().data

    rows = "".join(f"<tr><td>{it['product_name']}</td><td>{it['quantity']}</td><td>GHS {float(it['unit_cost']):,.2f}</td><td>GHS {float(it['line_total']):,.2f}</td></tr>" for it in items)

    body = f"""
    <div class="card">
        <h2>📦 Purchase Order {po['po_number']}</h2>
        <p><strong>Supplier:</strong> {po.get('supplier_name','')}</p>
        <p><strong>Date:</strong> {po.get('created_at','')[:16]}</p>
        <p><strong>Recorded by:</strong> {po.get('recorded_by','')}</p>
        <hr>
        <div class="table-wrap"><table>
            <tr><th>Item</th><th>Qty</th><th>Unit Cost</th><th>Total</th></tr>
            {rows}
        </table></div>
        <hr>
        <h3 style="text-align:right;">Total: GHS {float(po['total']):,.2f}</h3>
        <p style="text-align:right;">Paid: GHS {float(po.get('amount_paid', 0)):,.2f}</p>
        {f'<p style="text-align:right;color:#dc2626;">Owed: GHS {float(po.get("amount_owed", 0)):,.2f}</p>' if float(po.get("amount_owed", 0)) > 0 else ''}
        <p style="text-align:right;">Payment: {po.get('payment_method','')}</p>
        {f'<p><strong>Note:</strong> {po.get("note","")}</p>' if po.get("note") else ''}
    </div>
    <div class="card">
        <a href="/purchases" class="btn">← Back to Purchases</a>
        <a href="/suppliers/view/{po['supplier_id']}" class="btn">View Supplier</a>
    </div>
    """
    return HTMLResponse(content=page("Purchase", body, username, role))


# ============ EXPENSES ============

@app.get("/expenses", response_class=HTMLResponse)
def expenses_list(request: Request, month: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"
    if role != "admin":
        raise HTTPException(403, "Only admins can view expenses")

    expenses = supabase.table("expenses").select("*").order("expense_date", desc=True).execute().data

    now = datetime.now()
    if not month:
        month = now.strftime("%Y-%m")

    filtered = [e for e in expenses if str(e.get("expense_date", ""))[:7] == month]
    total_month = sum(float(e.get("amount", 0)) for e in filtered)

    cat_totals = {}
    for e in filtered:
        cat = e.get("category_name", "Unknown")
        cat_totals[cat] = cat_totals.get(cat, 0) + float(e.get("amount", 0))

    rows = "".join(f"""<tr>
        <td>{e.get('expense_date','')}</td>
        <td>{e.get('category_name','')}</td>
        <td>GHS {float(e.get('amount',0)):,.2f}</td>
        <td>{e.get('description','') or ''}</td>
        <td>{e.get('paid_to','') or ''}</td>
        <td>{e.get('payment_method','')}</td>
        <td><a href='/expenses/delete/{e['id']}' class='btn btn-danger btn-small' onclick="return confirm('Delete this expense?')">🗑️</a></td>
    </tr>""" for e in filtered)

    cat_rows = ""
    for cat, amt in sorted(cat_totals.items(), key=lambda x: -x[1]):
        pct = (amt / total_month * 100) if total_month > 0 else 0
        cat_rows += f"<tr><td>{cat}</td><td>GHS {amt:,.2f}</td><td>{pct:.0f}%</td></tr>"

    body = f"""
    <h2>💰 Expenses</h2>
    <div class="card">
        <a href="/expenses/new" class="btn btn-success">➕ Add Expense</a>
        <a href="/expenses/export?month={month}" class="btn">📥 Export Month to Excel</a>
    </div>
    <div class="card">
        <form method="get" style="display:flex;gap:10px;flex-wrap:wrap;align-items:flex-end;">
            <div style="flex:1;min-width:200px;">
                <label>Month (YYYY-MM)</label>
                <input type="text" name="month" value="{month}" placeholder="2026-09">
            </div>
            <button type="submit">Filter</button>
        </form>
    </div>
    <div class="card" style="background:#fef3c7;">
        <h3 style="margin:0;">Total for {month}: <span style="color:#dc2626;">GHS {total_month:,.2f}</span></h3>
        <p style="margin:5px 0 0 0;">{len(filtered)} expenses recorded</p>
    </div>
    <div class="card">
        <h3>📊 Breakdown by Category</h3>
        <div class="table-wrap"><table>
            <tr><th>Category</th><th>Amount</th><th>% of Total</th></tr>
            {cat_rows if cat_rows else "<tr><td colspan='3'>No expenses this month.</td></tr>"}
        </table></div>
    </div>
    <div class="card">
        <h3>📋 All Expenses ({len(filtered)})</h3>
        <div class="table-wrap"><table>
            <tr><th>Date</th><th>Category</th><th>Amount</th><th>Description</th><th>Paid To</th><th>Method</th><th></th></tr>
            {rows if rows else "<tr><td colspan='7'>No expenses this month.</td></tr>"}
        </table></div>
    </div>
    """
    return HTMLResponse(content=page("Expenses", body, username, role))


@app.get("/expenses/new", response_class=HTMLResponse)
def expense_new_form(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add expenses")

    cats = supabase.table("expense_categories").select("*").order("name").execute().data
    cat_options = "".join(f"<option value='{c['id']}' data-name='{c['name']}'>{c['name']}</option>" for c in cats)
    today = datetime.now().strftime("%Y-%m-%d")

    body = f"""
    <h2>➕ Add Expense</h2>
    <div class="card">
        <form method="post" action="/expenses/new">
            <label>Date</label>
            <input type="date" name="expense_date" value="{today}" required>
            <label>Category</label>
            <select name="category_id" required><option value="">-- Select category --</option>{cat_options}</select>
            <label>Amount (GHS)</label>
            <input type="number" step="0.01" name="amount" required min="0.01">
            <label>Description</label>
            <input type="text" name="description" placeholder="e.g. September rent">
            <label>Paid To</label>
            <input type="text" name="paid_to" placeholder="e.g. Landlord name">
            <label>Payment Method</label>
            <select name="payment_method">
                <option value="Cash">💵 Cash</option>
                <option value="Mobile Money">📱 Mobile Money</option>
                <option value="Bank Transfer">🏦 Bank Transfer</option>
                <option value="Card">💳 Card</option>
            </select>
            <button type="submit" class="btn btn-success">Save Expense</button>
            <a href="/expenses" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Add Expense", body, username, info.get("role")))


@app.post("/expenses/new")
async def expense_new(request: Request, expense_date: str = Form(...), category_id: str = Form(...), amount: float = Form(...), description: str = Form(""), paid_to: str = Form(""), payment_method: str = Form("Cash")):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add expenses")

    cat = supabase.table("expense_categories").select("*").eq("id", int(category_id)).single().execute().data
    cat_name = cat["name"] if cat else "Unknown"

    supabase.table("expenses").insert({"expense_date": expense_date, "category_id": int(category_id), "category_name": cat_name, "amount": amount, "description": description or None, "paid_to": paid_to or None, "payment_method": payment_method, "recorded_by": username}).execute()

    log(username, "expense_add", f"Added expense: '{cat_name}' GHS {amount:,.2f} - {description or 'no description'}", "info")
    return RedirectResponse("/expenses", status_code=303)


@app.get("/expenses/delete/{expense_id}")
def expense_delete(request: Request, expense_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can delete expenses")
    old = supabase.table("expenses").select("*").eq("id", expense_id).single().execute().data
    supabase.table("expenses").delete().eq("id", expense_id).execute()
    log(username, "expense_delete", f"DELETED expense: '{old.get('category_name','')}' GHS {float(old.get('amount',0)):,.2f}", "danger")
    return RedirectResponse("/expenses", status_code=303)


@app.get("/expenses/export")
def expenses_export(request: Request, month: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export expenses")

    if not month:
        month = datetime.now().strftime("%Y-%m")

    all_expenses = supabase.table("expenses").select("*").order("expense_date", desc=True).execute().data
    expenses = [e for e in all_expenses if str(e.get("expense_date", ""))[:7] == month]

    wb = Workbook()
    ws = wb.active
    ws.title = "Expenses"
    headers = ["Date", "Category", "Amount", "Description", "Paid To", "Payment Method", "Recorded By"]
    style_header(ws, headers)
    for e in expenses:
        ws.append([str(e.get("expense_date", "")), e.get("category_name", ""), float(e.get("amount", 0)), e.get("description", "") or "", e.get("paid_to", "") or "", e.get("payment_method", ""), e.get("recorded_by", "") or ""])

    widths = [15, 15, 12, 30, 20, 15, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_expenses_{month}.xlsx"}
    )


# ============ LOW STOCK ALERTS ============

@app.get("/alerts", response_class=HTMLResponse)
def alerts_page(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can view alerts")

    products = supabase.table("products").select("*").eq("is_active", True).execute().data
    low_stock = []
    out_of_stock = []
    for p in products:
        qty = float(p.get("quantity_in_stock", 0))
        reorder = float(p.get("reorder_level", 0))
        if qty <= 0:
            out_of_stock.append(p)
        elif qty <= reorder:
            low_stock.append(p)

    low_stock.sort(key=lambda p: float(p.get("quantity_in_stock", 0)))

    out_rows = "".join(f"<tr><td><strong>{p['name']}</strong><br><small>{p.get('sku','')}</small></td><td>{p.get('unit','')}</td><td class='low'>0</td><td>{p.get('reorder_level', 0)}</td><td><a href='/purchases/new' class='btn btn-success btn-small'>📦 Restock</a></td></tr>" for p in out_of_stock)
    low_rows = "".join(f"<tr><td><strong>{p['name']}</strong><br><small>{p.get('sku','')}</small></td><td>{p.get('unit','')}</td><td class='low'>{p.get('quantity_in_stock', 0)}</td><td>{p.get('reorder_level', 0)}</td><td><a href='/purchases/new' class='btn btn-success btn-small'>📦 Restock</a></td></tr>" for p in low_stock)

    body = f"""
    <h2>🚨 Low Stock Alerts</h2>
    <div class="card alert-box">
        <h3 style="color:#dc2626;margin:0;">⚠️ Summary</h3>
        <p style="font-size:18px;margin:10px 0 0 0;">
            <strong>{len(out_of_stock)}</strong> items OUT OF STOCK &nbsp;·&nbsp;
            <strong>{len(low_stock)}</strong> items running LOW
        </p>
    </div>
    <div class="card">
        <h3 style="color:#dc2626;">❌ Out of Stock ({len(out_of_stock)})</h3>
        <div class="table-wrap"><table>
            <tr><th>Material</th><th>Unit</th><th>In Stock</th><th>Reorder At</th><th></th></tr>
            {out_rows if out_rows else "<tr><td colspan='5'>None. Great!</td></tr>"}
        </table></div>
    </div>
    <div class="card">
        <h3 style="color:#d97706;">⚠️ Running Low ({len(low_stock)})</h3>
        <div class="table-wrap"><table>
            <tr><th>Material</th><th>Unit</th><th>In Stock</th><th>Reorder At</th><th></th></tr>
            {low_rows if low_rows else "<tr><td colspan='5'>None. Great!</td></tr>"}
        </table></div>
    </div>
    <div class="card">
        <a href="/products" class="btn">View All Materials</a>
        <a href="/" class="btn">Dashboard</a>
    </div>
    """
    return HTMLResponse(content=page("Alerts", body, username, info.get("role")))


# ============ DAILY SUMMARY ============

@app.get("/summary", response_class=HTMLResponse)
def daily_summary(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can view the summary")

    today = datetime.now().date()
    sales = supabase.table("sales").select("*").order("created_at", desc=True).execute().data

    today_sales = [s for s in sales if parse_dt(s.get("created_at", "")) and parse_dt(s.get("created_at", "")).date() == today]

    total_today = 0.0
    credit_today = 0.0
    payment_breakdown = {}
    cashier_breakdown = {}
    for s in today_sales:
        amount = float(s.get("total", 0))
        total_today += amount
        credit_today += float(s.get("amount_on_credit", 0))
        pm = s.get("payment_method", "Cash")
        payment_breakdown[pm] = payment_breakdown.get(pm, 0) + amount
        cashier = s.get("cashier_name") or s.get("user_id", "Unknown")
        cashier_breakdown[cashier] = cashier_breakdown.get(cashier, 0) + amount

    if today_sales:
        sale_ids = [s["id"] for s in today_sales]
        all_items = supabase.table("sale_items").select("*").execute().data
        today_items = [it for it in all_items if it.get("sale_id") in sale_ids]
    else:
        today_items = []

    product_totals = {}
    for it in today_items:
        name = it.get("product_name", "Unknown")
        qty = float(it.get("quantity", 0))
        rev = float(it.get("line_total", 0))
        if name not in product_totals:
            product_totals[name] = {"qty": 0, "revenue": 0}
        product_totals[name]["qty"] += qty
        product_totals[name]["revenue"] += rev

    top_today = sorted(product_totals.items(), key=lambda x: x[1]["qty"], reverse=True)[:5]
    top_rows = "".join(f"<tr><td>{name}</td><td>{data['qty']:.0f}</td><td>GHS {data['revenue']:,.2f}</td></tr>" for name, data in top_today)

    pay_rows = ""
    for method, amount in sorted(payment_breakdown.items(), key=lambda x: -x[1]):
        pct = (amount / total_today * 100) if total_today > 0 else 0
        pay_rows += f"<tr><td>{method}</td><td>GHS {amount:,.2f}</td><td>{pct:.0f}%</td></tr>"

    cashier_rows = "".join(f"<tr><td>{c}</td><td>GHS {amount:,.2f}</td></tr>" for c, amount in sorted(cashier_breakdown.items(), key=lambda x: -x[1]))

    body = f"""
    <h2>📅 Daily Summary — {today.strftime('%A, %d %B %Y')}</h2>
    <div class="card" style="background:#dbeafe;border-left:6px solid #1e40af;">
        <div class="big-num">GHS {total_today:,.2f}</div>
        <div style="text-align:center;color:#666;font-size:16px;">Total Sales Today</div>
    </div>
    <div class="grid">
        <div class="card stat"><div class="num">{len(today_sales)}</div><div class="label">Sales Made</div></div>
        <div class="card stat"><div class="num">{len(today_items)}</div><div class="label">Items Sold</div></div>
    </div>
    <div class="grid">
        <div class="card stat"><div class="num" style="color:#dc2626;">GHS {credit_today:,.2f}</div><div class="label">On Credit Today</div></div>
        <div class="card stat"><div class="num" style="color:#16a34a;">GHS {total_today - credit_today:,.2f}</div><div class="label">Cash/Paid Today</div></div>
    </div>
    <div class="card">
        <h3>💳 Payment Breakdown</h3>
        <div class="table-wrap"><table>
            <tr><th>Method</th><th>Amount</th><th>%</th></tr>
            {pay_rows if pay_rows else "<tr><td colspan='3'>No sales today yet.</td></tr>"}
        </table></div>
    </div>
    <div class="card">
        <h3>🏆 Top Materials Today</h3>
        <div class="table-wrap"><table>
            <tr><th>Material</th><th>Qty Sold</th><th>Revenue</th></tr>
            {top_rows if top_rows else "<tr><td colspan='3'>No materials sold today.</td></tr>"}
        </table></div>
    </div>
    <div class="card">
        <h3>👤 Sales by Cashier</h3>
        <div class="table-wrap"><table>
            <tr><th>Cashier</th><th>Amount</th></tr>
            {cashier_rows if cashier_rows else "<tr><td colspan='2'>No sales today yet.</td></tr>"}
        </table></div>
    </div>
    <div class="card" style="text-align:center;">
        <button onclick="window.print()" class="btn btn-success">🖨️ Print Summary</button>
        <a href="/" class="btn">Dashboard</a>
        <a href="/reports" class="btn">Full Reports</a>
    </div>
    """
    return HTMLResponse(content=page("Daily Summary", body, username, info.get("role")))


# ============ CART & SALES ============

@app.get("/sell", response_class=HTMLResponse)
def sell_page(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"
    products = supabase.table("products").select("*").eq("is_active", True).order("name").execute().data
    rows = ""
    for p in products:
        qty = float(p.get("quantity_in_stock", 0))
        rows += f"""<tr>
            <td><strong>{p['name']}</strong><br><small>{p.get('sku','')}</small></td>
            <td>{qty} {p.get('unit','')}</td>
            <td>GHS {float(p['selling_price']):,.2f}</td>
            <td>
                <form method="post" action="/cart/add" style="display:flex;gap:5px;align-items:center;flex-wrap:wrap;">
                    <input type="hidden" name="product_id" value="{p['id']}">
                    <input type="number" step="0.01" name="quantity" value="1" min="0.01" max="{qty}" style="width:70px;margin:0;padding:6px;">
                    <button type="submit" class="btn btn-success btn-small">+ Add</button>
                </form>
            </td>
        </tr>"""
    body = f"""
    <h2>New Sale — Add Items to Cart</h2>
    <div class="card">
        <a href="/scan" class="btn btn-success" style="margin-bottom:10px;">📷 Scan Barcode</a>
        <a href="/cart" class="btn btn-success" style="margin-bottom:10px;">🛒 View Cart & Checkout</a>
        <div class="table-wrap"><table>
            <tr><th>Material</th><th>Stock</th><th>Price</th><th>Add to Cart</th></tr>
            {rows if rows else "<tr><td colspan='4'>No materials.</td></tr>"}
        </table></div>
    </div>
    """
    return HTMLResponse(content=page("Sell", body, username, role))


@app.post("/cart/add")
async def cart_add(request: Request, product_id: int = Form(...), quantity: float = Form(...)):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    cart = get_cart(request)
    found = False
    for item in cart:
        if item["product_id"] == product_id:
            item["quantity"] += quantity
            found = True
            break
    if not found:
        cart.append({"product_id": product_id, "quantity": quantity})
    return save_cart(cart)


@app.get("/cart", response_class=HTMLResponse)
def cart_page(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"
    cart = get_cart(request)
    if not cart:
        body = """
        <h2>🛒 Your Cart</h2>
        <div class="card">
            <p>Your cart is empty.</p>
            <a href="/sell" class="btn">Start Adding Items</a>
            <a href="/scan" class="btn btn-success">📷 Scan Barcode</a>
        </div>
        """
        return HTMLResponse(content=page("Cart", body, username, role))
    ids = [c["product_id"] for c in cart]
    products = supabase.table("products").select("*").in_("id", ids).execute().data
    product_map = {p["id"]: p for p in products}
    rows = ""
    total = 0
    for item in cart:
        p = product_map.get(item["product_id"])
        if not p:
            continue
        line = float(p["selling_price"]) * item["quantity"]
        total += line
        rows += f"""<tr>
            <td>{p['name']}</td>
            <td>{item['quantity']} {p.get('unit','')}</td>
            <td>GHS {float(p['selling_price']):,.2f}</td>
            <td>GHS {line:,.2f}</td>
            <td>
                <form method="post" action="/cart/remove" style="display:inline;">
                    <input type="hidden" name="product_id" value="{p['id']}">
                    <button type="submit" class="btn btn-danger btn-small">Remove</button>
                </form>
            </td>
        </tr>"""

    customers = supabase.table("customers").select("*").eq("is_active", True).order("name").execute().data
    cust_options = "".join(f"<option value='{c['id']}'>{c['name']} — {c.get('phone','') or ''}</option>" for c in customers)

    body = f"""
    <h2>🛒 Your Cart</h2>
    <div class="card">
        <div class="table-wrap"><table>
            <tr><th>Material</th><th>Qty</th><th>Price</th><th>Subtotal</th><th></th></tr>
            {rows}
        </table></div>
        <div class="cart-total"><strong>Subtotal: GHS {total:,.2f}</strong></div>
        <br>
        <form method="post" action="/cart/checkout">
            <label>Customer (leave blank for walk-in)</label>
            <select name="customer_id">
                <option value="">— Walk-in customer —</option>
                {cust_options}
            </select>
            <label>Discount Type</label>
            <select name="discount_type">
                <option value="none">No Discount</option>
                <option value="percent">Percentage (%)</option>
                <option value="fixed">Fixed Amount (GHS)</option>
            </select>
            <label>Discount Value (0 if none)</label>
            <input type="number" step="0.01" name="discount_value" value="0" min="0">
            <label>Payment Method</label>
            <select name="payment_method">
                <option value="Cash">💵 Cash</option>
                <option value="Mobile Money">📱 Mobile Money</option>
                <option value="Bank Transfer">🏦 Bank Transfer</option>
                <option value="Card">💳 Card</option>
                <option value="Credit">📝 Credit (Customer Owes)</option>
            </select>
            <label>Amount Paid Now (if credit)</label>
            <input type="number" step="0.01" name="amount_paid_now" value="0" min="0">
            <br>
            <button type="submit" class="btn btn-success">✅ Complete Sale</button>
            <a href="/sell" class="btn">+ Add More</a>
            <a href="/scan" class="btn">📷 Scan More</a>
            <a href="/cart/clear" class="btn btn-danger">Clear Cart</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Cart", body, username, role))


@app.post("/cart/remove")
async def cart_remove(request: Request, product_id: int = Form(...)):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    cart = get_cart(request)
    cart = [item for item in cart if item["product_id"] != product_id]
    return save_cart(cart)


@app.get("/cart/clear")
def cart_clear(request: Request):
    response = RedirectResponse("/sell", status_code=303)
    response.delete_cookie("cart")
    return response


@app.post("/cart/checkout")
async def cart_checkout(request: Request, customer_id: str = Form(""), discount_type: str = Form("none"), discount_value: float = Form(0), payment_method: str = Form("Cash"), amount_paid_now: float = Form(0)):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    cart = get_cart(request)
    if not cart:
        return RedirectResponse("/sell", status_code=303)
    ids = [c["product_id"] for c in cart]
    products = supabase.table("products").select("*").in_("id", ids).execute().data
    product_map = {p["id"]: p for p in products}
    for item in cart:
        p = product_map.get(item["product_id"])
        if not p:
            continue
        if float(p["quantity_in_stock"]) < item["quantity"]:
            raise HTTPException(400, f"Not enough stock for {p['name']}")

    invoice_no = f"INV-{datetime.now().strftime('%Y%m%d%H%M%S')}"
    subtotal = 0.0
    for item in cart:
        p = product_map[item["product_id"]]
        subtotal += float(p["selling_price"]) * item["quantity"]

    discount_amount = 0.0
    if discount_type == "percent":
        discount_amount = subtotal * (discount_value / 100)
    elif discount_type == "fixed":
        discount_amount = discount_value
    if discount_amount > subtotal:
        discount_amount = subtotal
    final_total = subtotal - discount_amount

    customer_name = None
    cust_obj = None
    if customer_id:
        try:
            cust_obj = supabase.table("customers").select("*").eq("id", int(customer_id)).single().execute().data
            customer_name = cust_obj["name"] if cust_obj else None
        except Exception:
            cust_obj = None

    credit_amount = 0.0
    if payment_method == "Credit" and cust_obj:
        credit_amount = final_total - amount_paid_now
        if credit_amount < 0:
            credit_amount = 0
    elif payment_method == "Credit" and not cust_obj:
        raise HTTPException(400, "Credit sales require a registered customer")

    sale_data = {"invoice_no": invoice_no, "customer_id": int(customer_id) if customer_id else None, "customer_name": customer_name, "subtotal": subtotal, "discount": discount_amount, "total": final_total, "payment_method": payment_method, "amount_paid": final_total - credit_amount if payment_method == "Credit" else final_total, "amount_paid_now": amount_paid_now if payment_method == "Credit" else final_total, "amount_on_credit": credit_amount, "status": "completed", "user_id": username, "cashier_name": username}
    sale_result = supabase.table("sales").insert(sale_data).execute()
    sale_id = sale_result.data[0]["id"]

    for item in cart:
        p = product_map[item["product_id"]]
        qty = item["quantity"]
        unit_price = float(p["selling_price"])
        cost_price = float(p["cost_price"])
        line_total = qty * unit_price
        supabase.table("sale_items").insert({"sale_id": sale_id, "product_id": p["id"], "product_name": p["name"], "quantity": qty, "unit_price": unit_price, "cost_price": cost_price, "line_total": line_total}).execute()
        new_qty = float(p["quantity_in_stock"]) - qty
        supabase.table("products").update({"quantity_in_stock": new_qty}).eq("id", p["id"]).execute()
        supabase.table("stock_movements").insert({"product_id": p["id"], "movement_type": "OUT", "quantity": qty, "note": f"{invoice_no} — {p['name']} x {qty} — {username}"}).execute()

    if credit_amount > 0 and cust_obj:
        new_balance = float(cust_obj.get("balance", 0)) + credit_amount
        supabase.table("customers").update({"balance": new_balance}).eq("id", cust_obj["id"]).execute()

    cust_display = customer_name or "Walk-in"
    log(username, "sale", f"Sale {invoice_no} - Customer: {cust_display} - Total: GHS {final_total:,.2f} - Payment: {payment_method}", "info")

    response = RedirectResponse(f"/receipt/{sale_id}", status_code=303)
    response.delete_cookie("cart")
    return response


# ============ RECEIPT (WITH WHATSAPP) ============

@app.get("/receipt/{sale_id}", response_class=HTMLResponse)
def receipt(request: Request, sale_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    role = info.get("role", "cashier") if info else "cashier"
    sale = supabase.table("sales").select("*").eq("id", sale_id).single().execute().data
    items = supabase.table("sale_items").select("*").eq("sale_id", sale_id).execute().data

    rows = "".join(f"<tr><td>{it['product_name']}</td><td>{it['quantity']}</td><td>GHS {float(it['unit_price']):,.2f}</td><td>GHS {float(it['line_total']):,.2f}</td></tr>" for it in items)

    customer_line = ""
    if sale.get("customer_name"):
        customer_line = f"<p><strong>Customer:</strong> {sale['customer_name']}</p>"

    credit_line = ""
    if float(sale.get("amount_on_credit", 0)) > 0:
        credit_line = f'<p style="text-align:right;color:#dc2626;">On Credit: GHS {float(sale.get("amount_on_credit", 0)):,.2f}</p>'

    # Build plain text receipt for WhatsApp / SMS
    text_lines = [
        f"🏗️ {SHOP_NAME}",
        f"{SHOP_ADDRESS}",
        f"📞 {SHOP_PHONE}",
        "",
        f"Invoice: {sale['invoice_no']}",
        f"Date: {sale['created_at'][:16]}",
    ]
    if sale.get("customer_name"):
        text_lines.append(f"Customer: {sale['customer_name']}")
    text_lines.append("")
    text_lines.append("Items:")
    for it in items:
        text_lines.append(f"• {it['product_name']} — {it['quantity']} × GHS {float(it['unit_price']):,.2f} = GHS {float(it['line_total']):,.2f}")
    text_lines.append("")
    text_lines.append(f"Subtotal: GHS {float(sale['subtotal']):,.2f}")
    if float(sale.get("discount", 0)) > 0:
        text_lines.append(f"Discount: -GHS {float(sale.get('discount', 0)):,.2f}")
    text_lines.append(f"TOTAL: GHS {float(sale['total']):,.2f}")
    text_lines.append(f"Payment: {sale.get('payment_method', 'Cash')}")
    if float(sale.get("amount_on_credit", 0)) > 0:
        text_lines.append(f"On Credit: GHS {float(sale.get('amount_on_credit', 0)):,.2f}")
    text_lines.append("")
    text_lines.append("Thank you for your business! 🙏")

    text_receipt = "\n".join(text_lines)
    encoded_text = urllib.parse.quote(text_receipt)

    # Customer phone (if registered)
    customer_phone = ""
    if sale.get("customer_id"):
        try:
            c = supabase.table("customers").select("*").eq("id", sale["customer_id"]).single().execute().data
            if c and c.get("phone"):
                customer_phone = str(c["phone"]).replace(" ", "").replace("+", "").replace("-", "")
                if customer_phone.startswith("0"):
                    customer_phone = "233" + customer_phone[1:]
        except Exception:
            customer_phone = ""

    body = f"""
    <div class="card" id="receipt">
        <h2 style="text-align:center;">🏗️ {SHOP_NAME}</h2>
        <p style="text-align:center;">{SHOP_ADDRESS}</p>
        <p style="text-align:center;">📞 {SHOP_PHONE}</p>
        <hr>
        <p><strong>Invoice:</strong> {sale['invoice_no']}</p>
        <p><strong>Date:</strong> {sale['created_at'][:16]}</p>
        <p><strong>Cashier:</strong> {sale.get('cashier_name') or sale.get('user_id','')}</p>
        {customer_line}
        <hr>
        <div class="table-wrap"><table>
            <tr><th>Item</th><th>Qty</th><th>Price</th><th>Total</th></tr>
            {rows}
        </table></div>
        <hr>
        <p style="text-align:right;">Subtotal: GHS {float(sale['subtotal']):,.2f}</p>
        {f'<p style="text-align:right;">Discount: -GHS {float(sale.get("discount", 0)):,.2f}</p>' if float(sale.get("discount", 0)) > 0 else ''}
        <h3 style="text-align:right;">Total: GHS {float(sale['total']):,.2f}</h3>
        <p style="text-align:right;">Payment: {sale.get('payment_method','Cash')}</p>
        {credit_line}
        <hr>
        <p style="text-align:center;">Thank you for your business!</p>
    </div>

    <div class="card no-print" style="text-align:center;background:#dcfce7;border:2px solid #16a34a;">
        <h3 style="margin-top:0;color:#166534;">📤 Send Receipt to Customer</h3>
        <p style="color:#666;font-size:14px;">Choose how to send:</p>
        <a href="https://wa.me/{customer_phone}?text={encoded_text}" target="_blank" class="btn btn-success" style="background:#25D366;font-size:15px;padding:12px 20px;">📱 Send via WhatsApp</a>
        <a href="sms:{customer_phone}?body={encoded_text}" class="btn" style="font-size:15px;padding:12px 20px;">📤 Send via SMS</a>
        <button onclick="copyReceipt()" class="btn" style="font-size:15px;padding:12px 20px;background:#6b7280;">📋 Copy Receipt</button>
        {f'<p style="margin-top:10px;color:#666;font-size:13px;">Customer phone: +{customer_phone}</p>' if customer_phone else '<p style="margin-top:10px;color:#d97706;font-size:13px;">ℹ️ Customer phone not saved — WhatsApp will ask you to pick a contact</p>'}
    </div>

    <div class="card no-print" style="text-align:center;">
        <button onclick="window.print()" class="btn btn-success">🖨️ Print Receipt</button>
        <a href="/sell" class="btn">New Sale</a>
        <a href="/" class="btn">Dashboard</a>
    </div>

    <textarea id="receiptText" style="position:absolute;left:-9999px;">{text_receipt}</textarea>

    <script>
    function copyReceipt() {{
        var textArea = document.getElementById('receiptText');
        textArea.style.position = 'fixed';
        textArea.style.left = '0';
        textArea.style.top = '0';
        textArea.select();
        try {{
            document.execCommand('copy');
            alert('✅ Receipt copied! You can now paste it anywhere.');
        }} catch (err) {{
            alert('❌ Could not copy. Please try again.');
        }}
        textArea.style.position = 'absolute';
        textArea.style.left = '-9999px';
    }}
    </script>

    <style>
    @media print {{
        .header, .btn, button, .no-print {{ display: none !important; }}
        body {{ background: white; }}
        .card {{ box-shadow: none; }}
    }}
    </style>
    """
    return HTMLResponse(content=page("Receipt", body, username, role))


# ============ CATEGORIES ============

@app.get("/categories", response_class=HTMLResponse)
def categories_list(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can manage categories")
    cats = supabase.table("categories").select("*").order("name").execute().data
    products = supabase.table("products").select("*").eq("is_active", True).execute().data
    counts = {}
    for p in products:
        cid = p.get("category_id")
        counts[cid] = counts.get(cid, 0) + 1
    rows = "".join(f"<tr><td><strong>{c['name']}</strong></td><td>{c.get('description','')}</td><td>{counts.get(c['id'], 0)}</td><td><a href='/categories/delete/{c['id']}' class='btn btn-danger btn-small' onclick=\"return confirm('Delete {c['name']}?')\">🗑️</a></td></tr>" for c in cats)

    body = f"""
    <h2>Categories</h2>
    <div class="card">
        <h3>➕ Add New Category</h3>
        <form method="post" action="/categories/add">
            <label>Name</label>
            <input type="text" name="name" required placeholder="e.g. Roofing">
            <label>Description (optional)</label>
            <input type="text" name="description">
            <button type="submit" class="btn btn-success">Add Category</button>
        </form>
    </div>
    <div class="card">
        <h3>All Categories ({len(cats)})</h3>
        <div class="table-wrap"><table>
            <tr><th>Category</th><th>Description</th><th>Materials</th><th></th></tr>
            {rows if rows else "<tr><td colspan='4'>No categories yet.</td></tr>"}
        </table></div>
    </div>
    """
    return HTMLResponse(content=page("Categories", body, username, info.get("role")))


@app.post("/categories/add")
async def category_add(request: Request, name: str = Form(...), description: str = Form("")):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add categories")
    existing = supabase.table("categories").select("*").eq("name", name).execute().data
    if not existing:
        supabase.table("categories").insert({"name": name, "description": description or None}).execute()
    return RedirectResponse("/categories", status_code=303)


@app.get("/categories/delete/{category_id}")
def category_delete(request: Request, category_id: int):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can delete categories")
    others = supabase.table("categories").select("*").eq("name", "Others").execute().data
    if others:
        others_id = others[0]["id"]
    else:
        new_others = supabase.table("categories").insert({"name": "Others", "description": "Default"}).execute()
        others_id = new_others.data[0]["id"]
    supabase.table("products").update({"category_id": others_id}).eq("category_id", category_id).execute()
    supabase.table("categories").delete().eq("id", category_id).execute()
    log(username, "category_delete", f"Deleted category (moved materials to 'Others')", "warning")
    return RedirectResponse("/categories", status_code=303)


# ============ USERS ============

@app.get("/users", response_class=HTMLResponse)
def users_list(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can manage users")
    users = supabase.table("shop_users").select("*").order("username").execute().data
    rows = ""
    for u in users:
        role_badge = "🔴 Admin" if u.get("role") == "admin" else "🟢 Cashier"
        status_badge = "✅ Active" if u.get("is_active") else "❌ Inactive"
        actions = f"<a href='/users/edit/{u['id']}' class='btn btn-warn btn-small'>✏️</a>"
        if u["username"] != username:
            actions += f" <a href='/users/delete/{u['id']}' class='btn btn-danger btn-small' onclick=\"return confirm('Delete {u['username']}?')\">🗑️</a>"
        rows += f"<tr><td><strong>{u['username']}</strong><br><small>{u.get('full_name','')}</small></td><td>{role_badge}</td><td>{status_badge}</td><td>{actions}</td></tr>"

    body = f"""
    <h2>👥 Users</h2>
    <div class="card"><a href="/users/add" class="btn btn-success">➕ Add New User</a></div>
    <div class="card">
        <div class="table-wrap"><table>
            <tr><th>User</th><th>Role</th><th>Status</th><th>Actions</th></tr>
            {rows}
        </table></div>
    </div>
    """
    return HTMLResponse(content=page("Users", body, username, info.get("role")))


@app.get("/users/add", response_class=HTMLResponse)
def users_add_form(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can manage users")
    body = """
    <h2>➕ Add New User</h2>
    <div class="card">
        <form method="post" action="/users/add">
            <label>Username (login name)</label>
            <input type="text" name="username" required>
            <label>Password</label>
            <input type="text" name="password" required minlength="4">
            <label>Full Name</label>
            <input type="text" name="full_name">
            <label>Role</label>
            <select name="role" required>
                <option value="cashier">Cashier (limited access)</option>
                <option value="admin">Admin (full access)</option>
            </select>
            <button type="submit" class="btn btn-success">Create User</button>
            <a href="/users" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Add User", body, username, info.get("role")))


@app.post("/users/add")
async def users_add(request: Request, username: str = Form(...), password: str = Form(...), full_name: str = Form(""), role: str = Form("cashier")):
    current = get_current_user(request)
    if not current:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(current)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can add users")
    existing = supabase.table("shop_users").select("*").eq("username", username).execute().data
    if not existing:
        supabase.table("shop_users").insert({"username": username, "password": password, "full_name": full_name, "role": role, "is_active": True}).execute()
        log(current, "user_add", f"Created new user '{username}' with role '{role}'", "warning")
    return RedirectResponse("/users", status_code=303)


@app.get("/users/edit/{user_id}", response_class=HTMLResponse)
def users_edit_form(request: Request, user_id: int):
    current = get_current_user(request)
    if not current:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(current)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can manage users")
    u = supabase.table("shop_users").select("*").eq("id", user_id).single().execute().data
    body = f"""
    <h2>✏️ Edit User: {u['username']}</h2>
    <div class="card">
        <form method="post" action="/users/edit/{user_id}">
            <label>Full Name</label>
            <input type="text" name="full_name" value="{u.get('full_name','')}">
            <label>Password (leave empty to keep current)</label>
            <input type="text" name="password" placeholder="New password (or leave blank)">
            <label>Role</label>
            <select name="role">
                <option value="cashier" {'selected' if u.get('role')=='cashier' else ''}>Cashier</option>
                <option value="admin" {'selected' if u.get('role')=='admin' else ''}>Admin</option>
            </select>
            <label>Status</label>
            <select name="is_active">
                <option value="true" {'selected' if u.get('is_active') else ''}>Active</option>
                <option value="false" {'selected' if not u.get('is_active') else ''}>Inactive (blocked)</option>
            </select>
            <button type="submit" class="btn btn-success">Save Changes</button>
            <a href="/users" class="btn">Cancel</a>
        </form>
    </div>
    """
    return HTMLResponse(content=page("Edit User", body, current, info.get("role")))


@app.post("/users/edit/{user_id}")
async def users_edit(request: Request, user_id: int, full_name: str = Form(""), password: str = Form(""), role: str = Form("cashier"), is_active: str = Form("true")):
    current = get_current_user(request)
    if not current:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(current)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can edit users")
    data = {"full_name": full_name, "role": role, "is_active": (is_active == "true")}
    if password:
        data["password"] = password
    supabase.table("shop_users").update(data).eq("id", user_id).execute()
    log(current, "user_edit", f"Edited user ID {user_id} (role: {role}, active: {is_active})", "warning")
    return RedirectResponse("/users", status_code=303)


@app.get("/users/delete/{user_id}")
def users_delete(request: Request, user_id: int):
    current = get_current_user(request)
    if not current:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(current)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can delete users")
    u = supabase.table("shop_users").select("*").eq("id", user_id).single().execute().data
    if u["username"] == current:
        raise HTTPException(400, "You cannot delete your own account")
    supabase.table("shop_users").delete().eq("id", user_id).execute()
    log(current, "user_delete", f"DELETED user '{u.get('username','')}'", "danger")
    return RedirectResponse("/users", status_code=303)


# ============ REPORTS ============

@app.get("/reports", response_class=HTMLResponse)
def reports(request: Request, start: str = "", end: str = "", preset: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can view reports")

    now = datetime.now()
    today_str = now.strftime("%Y-%m-%d")

    if preset == "today":
        start = today_str
        end = today_str
    elif preset == "week":
        start = (now - timedelta(days=7)).strftime("%Y-%m-%d")
        end = today_str
    elif preset == "month":
        start = now.replace(day=1).strftime("%Y-%m-%d")
        end = today_str
    elif preset == "lastmonth":
        first_of_this_month = now.replace(day=1)
        last_of_prev = first_of_this_month - timedelta(days=1)
        start = last_of_prev.replace(day=1).strftime("%Y-%m-%d")
        end = last_of_prev.strftime("%Y-%m-%d")
    elif preset == "year":
        start = now.replace(month=1, day=1).strftime("%Y-%m-%d")
        end = today_str
    elif preset == "all":
        start = ""
        end = ""

    if not start:
        start = now.replace(day=1).strftime("%Y-%m-%d")
    if not end:
        end = today_str

    period_label = f"{start} to {end}"

    def in_range(date_str):
        if not date_str:
            return False
        d = str(date_str)[:10]
        return start <= d <= end

    all_sales = supabase.table("sales").select("*").order("created_at", desc=True).execute().data
    all_items = supabase.table("sale_items").select("*").execute().data
    all_expenses = supabase.table("expenses").select("*").execute().data

    period_sales = [s for s in all_sales if in_range(s.get("created_at", ""))]
    sale_ids = [s["id"] for s in period_sales]
    period_items = [it for it in all_items if it.get("sale_id") in sale_ids]
    period_expenses_list = [e for e in all_expenses if start <= str(e.get("expense_date", ""))[:10] <= end]
    period_expenses = sum(float(e.get("amount", 0)) for e in period_expenses_list)

    period_total = sum(float(s.get("total", 0)) for s in period_sales)
    period_credit = sum(float(s.get("amount_on_credit", 0)) for s in period_sales)
    period_revenue = sum(float(it.get("line_total", 0)) for it in period_items)
    period_cost = sum(float(it.get("cost_price", 0)) * float(it.get("quantity", 0)) for it in period_items)
    period_gross_profit = period_revenue - period_cost
    period_net_profit = period_gross_profit - period_expenses

    product_stats = {}
    for it in period_items:
        name = it.get("product_name", "Unknown")
        qty = float(it.get("quantity", 0))
        revenue = float(it.get("line_total", 0))
        cost = float(it.get("cost_price", 0)) * qty
        profit = revenue - cost
        if name not in product_stats:
            product_stats[name] = {"qty": 0, "revenue": 0, "profit": 0}
        product_stats[name]["qty"] += qty
        product_stats[name]["revenue"] += revenue
        product_stats[name]["profit"] += profit

    top = sorted(product_stats.items(), key=lambda x: x[1]["qty"], reverse=True)[:10]
    top_rows = "".join(f"<tr><td>{name}</td><td>{data['qty']:.0f}</td><td>GHS {data['revenue']:,.2f}</td><td>GHS {data['profit']:,.2f}</td></tr>" for name, data in top)

    recent_rows = ""
    for s in period_sales[:20]:
        cust = s.get("customer_name") or "Walk-in"
        recent_rows += f"<tr><td>{s.get('created_at','')[:16]}</td><td>{s.get('invoice_no','')}</td><td>{cust}</td><td>GHS {float(s.get('total',0)):,.2f}</td><td>{s.get('cashier_name') or s.get('user_id','')}</td><td><a href='/receipt/{s['id']}' class='btn btn-small'>View</a></td></tr>"

    body = f"""
    <h2>Reports</h2>
    <div class="card date-bar">
        <h3 style="margin-top:0;">📅 Select Period</h3>
        <form method="get" action="/reports">
            <div class="field"><label>From</label><input type="date" name="start" value="{start}"></div>
            <div class="field"><label>To</label><input type="date" name="end" value="{end}"></div>
            <button type="submit" class="btn">Apply</button>
        </form>
        <div class="quick-links">
            <a href="/reports?preset=today" class="btn btn-quick">Today</a>
            <a href="/reports?preset=week" class="btn btn-quick">Last 7 days</a>
            <a href="/reports?preset=month" class="btn btn-quick">This Month</a>
            <a href="/reports?preset=lastmonth" class="btn btn-quick">Last Month</a>
            <a href="/reports?preset=year" class="btn btn-quick">This Year</a>
            <a href="/reports?preset=all" class="btn btn-quick">All Time</a>
        </div>
    </div>
    <div class="card">
        <span class="period-badge">📅 Period: {period_label}</span>
        <h3>📥 Export to Excel</h3>
        <a href="/export/sales?start={start}&end={end}" class="btn btn-success">📥 Sales for this period</a>
        <a href="/export/products" class="btn btn-success">📥 All Materials</a>
        <a href="/export/stock" class="btn btn-success">📥 Stock Movements</a>
        <a href="/export/customers" class="btn btn-success">📥 Customers</a>
        <a href="/export/purchases" class="btn btn-success">📥 Purchases</a>
    </div>
    <div class="grid">
        <div class="card stat"><div class="num">{len(period_sales)}</div><div class="label">Sales in Period</div></div>
        <div class="card stat"><div class="num">GHS {period_total:,.2f}</div><div class="label">Total Revenue</div></div>
    </div>
    <div class="grid">
        <div class="card stat"><div class="num" style="color:#dc2626;">GHS {period_credit:,.2f}</div><div class="label">On Credit</div></div>
        <div class="card stat"><div class="num" style="color:#dc2626;">GHS {period_expenses:,.2f}</div><div class="label">Expenses</div></div>
    </div>
    <div class="grid">
        <div class="card stat"><div class="num" style="color:#16a34a;">GHS {period_gross_profit:,.2f}</div><div class="label">Gross Profit</div></div>
        <div class="card stat"><div class="num" style="color:{'#16a34a' if period_net_profit >= 0 else '#dc2626'};">GHS {period_net_profit:,.2f}</div><div class="label">Net Profit</div></div>
    </div>
    <div class="card {'profit-box' if period_net_profit >= 0 else 'loss-box'}">
        <h3 style="margin:0;">{'💰' if period_net_profit >= 0 else '⚠️'} {'Net Profit' if period_net_profit >= 0 else 'Net Loss'} for Period</h3>
        <div class="big-num" style="color:{'#16a34a' if period_net_profit >= 0 else '#dc2626'};">GHS {period_net_profit:,.2f}</div>
        <p style="text-align:center;color:#666;">Revenue − Cost of Goods − Expenses</p>
    </div>
    <div class="card" style="text-align:center;">
        <a href="/pnl?start={start}&end={end}" class="btn btn-success" style="font-size:16px;padding:14px 24px;">📊 View Full Profit & Loss Statement</a>
    </div>
    <div class="card">
        <h3>🏆 Top Selling Materials in Period</h3>
        <div class="table-wrap"><table>
            <tr><th>Material</th><th>Qty Sold</th><th>Revenue</th><th>Profit</th></tr>
            {top_rows if top_rows else "<tr><td colspan='4'>No sales in this period.</td></tr>"}
        </table></div>
    </div>
    <div class="card">
        <h3>🧾 Sales in Period ({len(period_sales)})</h3>
        <div class="table-wrap"><table>
            <tr><th>Date</th><th>Invoice</th><th>Customer</th><th>Total</th><th>Cashier</th><th></th></tr>
            {recent_rows if recent_rows else "<tr><td colspan='6'>No sales in this period.</td></tr>"}
        </table></div>
    </div>
    """
    return HTMLResponse(content=page("Reports", body, username, info.get("role")))


# ============ EXCEL EXPORTS ============

def style_header(ws, headers):
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")


@app.get("/export/sales")
def export_sales(request: Request, start: str = "", end: str = ""):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export data")

    all_sales = supabase.table("sales").select("*").order("created_at", desc=True).execute().data
    if start and end:
        sales = [s for s in all_sales if start <= str(s.get("created_at", ""))[:10] <= end]
        fname_period = f"{start}_to_{end}"
    else:
        sales = all_sales
        fname_period = datetime.now().strftime('%Y%m%d_%H%M')

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    headers = ["Date", "Invoice", "Customer", "Subtotal", "Discount", "Total", "Paid", "On Credit", "Payment", "Cashier"]
    style_header(ws, headers)
    for s in sales:
        ws.append([s.get("created_at", "")[:19].replace("T", " "), s.get("invoice_no", ""), s.get("customer_name") or "Walk-in", float(s.get("subtotal", 0)), float(s.get("discount", 0)), float(s.get("total", 0)), float(s.get("amount_paid_now", s.get("total", 0))), float(s.get("amount_on_credit", 0)), s.get("payment_method", ""), s.get("cashier_name") or s.get("user_id", "")])

    widths = [20, 22, 20, 12, 12, 12, 12, 12, 15, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_sales_{fname_period}.xlsx"}
    )


@app.get("/export/products")
def export_products(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export data")
    products = supabase.table("products").select("*").eq("is_active", True).order("name").execute().data
    wb = Workbook()
    ws = wb.active
    ws.title = "Materials"
    headers = ["Name", "SKU", "Unit", "Stock", "Cost Price", "Selling Price", "Reorder Level", "Location"]
    style_header(ws, headers)
    for p in products:
        ws.append([p.get("name", ""), p.get("sku", ""), p.get("unit", ""), float(p.get("quantity_in_stock", 0)), float(p.get("cost_price", 0)), float(p.get("selling_price", 0)), float(p.get("reorder_level", 0)), p.get("location") or ""])
    widths = [30, 15, 10, 10, 12, 14, 14, 15]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_materials_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"}
    )


@app.get("/export/stock")
def export_stock(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export data")
    movements = supabase.table("stock_movements").select("*").order("created_at", desc=True).limit(5000).execute().data
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock Movements"
    headers = ["Date", "Type", "Quantity", "Note"]
    style_header(ws, headers)
    for m in movements:
        ws.append([m.get("created_at", "")[:19].replace("T", " "), m.get("movement_type", ""), float(m.get("quantity", 0)), m.get("note", "")])
    widths = [20, 12, 12, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_stock_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"}
    )


@app.get("/export/customers")
def export_customers(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export data")
    customers = supabase.table("customers").select("*").eq("is_active", True).order("name").execute().data
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    headers = ["Name", "Phone", "Address", "Balance Owed", "Notes"]
    style_header(ws, headers)
    for c in customers:
        ws.append([c.get("name", ""), c.get("phone", "") or "", c.get("address", "") or "", float(c.get("balance", 0)), c.get("notes", "") or ""])
    widths = [25, 15, 25, 15, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_customers_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"}
    )


@app.get("/export/purchases")
def export_purchases(request: Request):
    username = get_current_user(request)
    if not username:
        return RedirectResponse("/login", status_code=303)
    info = get_user_info(username)
    if not info or info.get("role") != "admin":
        raise HTTPException(403, "Only admins can export data")
    purchases = supabase.table("purchase_orders").select("*").order("created_at", desc=True).execute().data
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchases"
    headers = ["Date", "PO Number", "Supplier", "Total", "Paid", "Owed", "Payment Method"]
    style_header(ws, headers)
    for po in purchases:
        ws.append([po.get("created_at", "")[:19].replace("T", " "), po.get("po_number", ""), po.get("supplier_name", ""), float(po.get("total", 0)), float(po.get("amount_paid", 0)), float(po.get("amount_owed", 0)), po.get("payment_method", "")])
    widths = [20, 22, 25, 12, 12, 12, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=OBOLO_purchases_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"}
    )


@app.get("/health")
def health():
    return {"status": "ok"}